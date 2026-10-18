"""Legacy `.doc` conversion: detection, backend selection, E4151 and LO round-trip."""
from __future__ import annotations

import subprocess
from pathlib import Path

import pytest

from doc2md.errors import E4151_UnsupportedType
from doc2md.pipeline import legacy_doc
from doc2md.pipeline.s1_parse import read_any, sniff

pytestmark = [pytest.mark.milestone("DoD")]
OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _fake_doc(tmp_path: Path, name: str = "legacy.doc") -> Path:
    p = tmp_path / name
    p.write_bytes(OLE + b"\x00" * 128)
    return p


def test_sniff_detects_legacy_doc_by_ole_magic(tmp_path):
    p = _fake_doc(tmp_path)
    assert sniff(p) == "doc"


def test_ole_with_non_doc_suffix_is_rejected(tmp_path):
    p = _fake_doc(tmp_path, "legacy.bin")
    with pytest.raises(E4151_UnsupportedType):
        sniff(p)


def test_backend_prefers_libreoffice(monkeypatch):
    monkeypatch.setattr(legacy_doc, "find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(legacy_doc, "find_word", lambda: True)
    assert legacy_doc.available_backend("auto") == "libreoffice"
    assert legacy_doc.available_backend("word") == "word"
    assert legacy_doc.available_backend("none") is None


def test_missing_backend_raises_e4151(tmp_path):
    p = _fake_doc(tmp_path)
    with pytest.raises(E4151_UnsupportedType) as ei:
        read_any(p, legacy_doc_backend="none")
    assert ei.value.detail.get("stage") == "backend_unavailable"


def test_libreoffice_command_uses_isolated_profile(tmp_path):
    cmd = legacy_doc._libreoffice_cmd(
        "soffice", tmp_path / "in.doc", tmp_path / "out", tmp_path / "profile")
    assert cmd[0] == "soffice"
    assert "--headless" in cmd and "--norestore" in cmd
    assert "--convert-to" in cmd and "docx" in cmd
    assert any(x.startswith("-env:UserInstallation=") for x in cmd)
    assert str(tmp_path / "in.doc") in cmd


@pytest.mark.skipif(legacy_doc.find_libreoffice() is None,
                    reason="LibreOffice not installed")
def test_libreoffice_doc_roundtrip(tmp_path):
    import docx

    src_docx = tmp_path / "sample.docx"
    d = docx.Document()
    d.add_heading("旧格式转换测试", level=1)
    d.add_paragraph("这是 DOC 转换后的正文段落。")
    d.save(str(src_docx))
    lo = legacy_doc.find_libreoffice()
    profile = tmp_path / "lo-profile"
    subprocess.run(
        [lo, "--headless", "--norestore",
         f"-env:UserInstallation={profile.as_uri()}",
         "--convert-to", "doc", "--outdir", str(tmp_path), str(src_docx)],
        check=True, capture_output=True, timeout=120)
    doc = tmp_path / "sample.doc"
    assert doc.is_file() and doc.read_bytes().startswith(OLE)

    res = read_any(doc)
    text = "\n".join(ln.text for pg in res for ln in pg.lines)
    assert "旧格式转换测试" in text
    assert "这是 DOC 转换后的正文段落。" in text
