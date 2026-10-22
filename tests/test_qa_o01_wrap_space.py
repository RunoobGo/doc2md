"""tests/test_qa_o01_wrap_space.py — O-01 折行空白保真（QA-D2M-2026-10-21 / C-04）。

缺陷（C-04）：`extract_text_lines` 的行字典把行缘空白从 text/chars/bbox 一并
剥掉，折行恰好切在空格上时（原文 "数字 2026" → 行 "…数字" + 行 "2026…"），
S3 的中西文拼接规则默认不插空格，空格被静默吃掉（语料 pdf_09 实测 1/22 处，
转换结果 "数字2026"）。

修复：
- S1 从 page.chars 补行缘空白几何证据——`Line.trailing_space` /
  `leading_space`（判据 `_edge_spaces`：只认与行缘相邻的空白字符；CID 字体
  未映射字形的重叠噪声空白与远距（栏间）空白均不算）；
- S3 的 `_join_line` / `_append_line` / `_flush` 与三个并回助手
  （`_join_back` / `_join_list_continuation` / `_join_heading_continuation`）
  在任一侧带证据时补回空格；块尾证据经 `Block.line_trail` 瞬态传递。

本文件锁定：位置判据的三个方向（相邻算 / 重叠噪声不算 / 远距不算）、拼接
语义（证据补空格、无证据不插、数字切开与连字符断词不受影响）、端到端
（单栏行尾/行首证据、双栏 crop 路径、DOCX 段落旗标）与序列化纪律
（line_trail 不进 ir.json）。
"""
from __future__ import annotations

import pytest

from doc2md.ir import Block
from doc2md.pipeline.s1_parse import Line, RawPage, _edge_spaces, read_pdf
from doc2md.pipeline.s3_structure import _append_line, _join_line, analyze

pytestmark = [pytest.mark.milestone("DoD")]


# ---------------------------------------------------------------- 工具
def _page(*lines: Line, index: int = 0) -> RawPage:
    return RawPage(index=index, lines=list(lines))


def _line(text: str, **kw) -> Line:
    return Line(text=text, **kw)


def _ch(text: str, x0: float, x1: float, top: float = 100.0) -> dict:
    return {"text": text, "x0": x0, "x1": x1, "top": top, "bottom": top + 10.0}


# ------------------------------------------------ S1：_edge_spaces 位置判据
def test_edge_spaces_detects_adjacent_blanks():
    """行缘相邻的空白字符 → leading / trailing 证据。"""
    vis = [_ch("数", 100.0, 110.0), _ch("2", 110.0, 118.0)]
    blanks = [_ch(" ", 96.0, 98.0), _ch(" ", 120.0, 122.4)]
    lead, trail = _edge_spaces(list(vis), list(blanks))
    assert lead is True and trail is True


def test_edge_spaces_rejects_midline_and_overlapping_blanks():
    """行内空格与压在可见字形上的空白都不算证据。

    语料 2537 个空白的实测形态：绝大多数是行内空格（两个可见字符之间），
    其位置必被相邻性判据排除；本用例同时锁定「与行缘重叠 > 0.5pt」的
    空白（宽度量级异常的空格/未映射字形）同样不算。
    """
    vis = [_ch("数", 100.0, 110.0), _ch("2", 110.0, 118.0)]
    midline = [_ch(" ", 112.0, 114.28)]                 # 行内空格（真实主形态）
    lead, trail = _edge_spaces(list(vis), list(midline))
    assert lead is False and trail is False
    overlap = [_ch(" ", 99.0, 100.86), _ch(" ", 109.5, 111.36)]
    lead, trail = _edge_spaces(list(vis), list(overlap))
    assert lead is False and trail is False


def test_edge_spaces_rejects_wide_far_blanks():
    """整体落在行缘 gap 之外的空白不算（双栏页同 band 内他栏空白靠此收口）；
    宽度超过 gap 的空白同样不算（真实空格宽度 ≈ 同字号字宽量级）。"""
    vis = [_ch("数", 100.0, 110.0), _ch("2", 110.0, 118.0)]
    far = [_ch(" ", 60.0, 62.4), _ch(" ", 160.0, 162.4)]
    lead, trail = _edge_spaces(list(vis), list(far))
    assert lead is False and trail is False
    wide = [_ch(" ", 90.0, 99.6), _ch(" ", 118.4, 130.0)]   # 宽 9.6 / 11.6 > gap 7.5
    lead, trail = _edge_spaces(list(vis), list(wide))
    assert lead is False and trail is False


def test_edge_spaces_empty_inputs():
    assert _edge_spaces([], [_ch(" ", 0.0, 2.0)]) == (False, False)
    assert _edge_spaces([_ch("数", 0.0, 10.0)], []) == (False, False)


# ------------------------------------------------ S3：拼接语义
def test_join_line_space_evidence_at_cjk_boundary():
    """中西文边界：无证据不插（中文排版惯例），有证据必插（原文确有空格）。"""
    assert _join_line("数字", "2026") == "2026"                    # 无证据：直连
    assert _join_line("数字", "2026", space=True) == " 2026"       # 有证据：补回
    assert _join_line("2026", "和标点", space=True) == " 和标点"   # 数字→中文
    assert _join_line("中文与", "English", space=True) == " English"
    assert _join_line("English", "中文", space=True) == " 中文"
    # 非 CJK 边界本就有空格规则，evidence 不改变既有口径
    assert _join_line("hello", "world") == " world"
    assert _join_line("hello", "world", space=True) == " world"


def test_append_line_space_evidence_end_to_end():
    assert _append_line("正文包含中文与数字", "2026 和标点", space=True) \
        == "正文包含中文与数字 2026 和标点"
    assert _append_line("正文包含中文与数字", "2026 和标点") \
        == "正文包含中文与数字2026 和标点"


def test_space_evidence_does_not_affect_digit_split_and_hyphen():
    """两条早退分支不受证据影响：折行切开同一个数字（"2"+"026"）与连字符
    断词（"under-"+"score"）在原文里都不可能有空格。"""
    assert _join_line("2", "026", space=True) == "026"
    assert _append_line("正文包含数字 2", "026 和标点。", space=True) \
        == "正文包含数字 2026 和标点。"
    assert _append_line("under-", "score", space=True) == "underscore"


# ------------------------------------------------ S3：analyze 端到端
def _wrap_ir(*lines: Line):
    return analyze([_page(*lines)])


def test_analyze_restores_trailing_space_evidence():
    """行尾证据（空格画在本行行尾）："…数字 " + "2026…" → "…数字 2026…"。"""
    ir = _wrap_ir(
        _line("这一段用于验证折行空白，正文包含中文与数字", size=10.5,
              trailing_space=True),
        _line("2026 和标点，转换后必须保留原有空格。", size=10.5),
    )
    assert [b.type for b in ir.blocks] == ["paragraph"]
    assert ir.blocks[0].text == ("这一段用于验证折行空白，正文包含中文与数字 "
                                 "2026 和标点，转换后必须保留原有空格。")


def test_analyze_restores_leading_space_evidence():
    """行首证据（空格画在下一行行首，pdf_09 的实证形态）。"""
    ir = _wrap_ir(
        _line("这一段用于验证折行空白，正文包含中文与数字", size=10.5),
        _line("2026 和标点，转换后必须保留原有空格。", size=10.5,
              leading_space=True),
    )
    assert ir.blocks[0].text == ("这一段用于验证折行空白，正文包含中文与数字 "
                                 "2026 和标点，转换后必须保留原有空格。")


def test_analyze_no_evidence_keeps_cjk_rule():
    """反例：原文折行处无空格（无 blank 字符）→ 不得插入（宁漏勿误）。"""
    ir = _wrap_ir(
        _line("这一段用于验证折行空白，正文包含中文与数字", size=10.5),
        _line("2026 和标点，转换后必须保留原有空格。", size=10.5),
    )
    assert ir.blocks[0].text == ("这一段用于验证折行空白，正文包含中文与数字"
                                 "2026 和标点，转换后必须保留原有空格。")


def test_analyze_join_back_uses_block_line_trail():
    """CH-430 续句并回路径：块尾证据经 Block.line_trail 传递。

    段落块由多行缝合而成时，块尾证据取**最后一行**的 trailing_space
    （`_flush` 收尾时写入）；续句标点起头且段缓冲已空（空行冲刷过）时
    走 `_join_back`，补空格判断 = 续行 leading_space or 块尾 line_trail。
    """
    ir = _wrap_ir(
        _line("第一行说明文字在这里结束", size=10.5, trailing_space=True),
        _line("第二行继续同一段落的正文", size=10.5, trailing_space=True),
        _line("", size=10.5),                      # 空行：冲刷段缓冲落块
        _line("，续句标点起头并回上一块。", size=10.5),
    )
    assert [b.type for b in ir.blocks] == ["paragraph"]
    assert ir.blocks[0].text == ("第一行说明文字在这里结束 第二行继续同一段落的正文"
                                 " ，续句标点起头并回上一块。")


# ------------------------------------------------ 双栏 crop 路径
def test_two_column_wrap_space_restored(tmp_path):
    """双栏路径：行字典 chars 被 extract_text_lines 剥掉空白，证据须由 crop
    后的 chars 补（`_extract_two_column` 的 edge_blanks_for）。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    CJK = "STSong-Light"
    path = tmp_path / "two_col_wrap.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    left = ["左栏折行空白测试：中文与数字 ", "2026 之间原有空格，",
            "以及中文与数字", " 2026 的又一次折行；", "左栏后续文字用于占位，",
            "继续写满多行以满足", "双栏检测的行数要求，", "末行文字与本栏平行。"]
    right = ["右栏对应文字：验证阅读顺序", "与栏间分隔，右栏包含中文",
             "和数字 2026 的内容；", "右栏后续文字用于占位，", "继续写满多行以满足",
             "双栏检测的行数要求，", "末行文字与对栏平行。"]
    for i, t in enumerate(left):
        c.setFont(CJK, 10)
        c.drawString(1.2 * cm, 842 - 3.0 * cm - i * 0.6 * cm, t)
    for i, t in enumerate(right):
        c.setFont(CJK, 10)
        c.drawString(12.0 * cm, 842 - 3.0 * cm - i * 0.6 * cm, t)
    c.showPage()
    c.save()

    pages = read_pdf(path)
    lines = pages[0].lines
    texts = [ln.text for ln in lines]
    # 双栏阅读顺序：左栏全部行在右栏之前（末行措辞避开「右栏/左栏」子串互串）
    left_idx = [i for i, t in enumerate(texts)
                if "左栏" in t or "本栏" in t or t.startswith("2026")
                or "以及中文与数字" in t]
    right_idx = [i for i, t in enumerate(texts) if "右栏" in t or "对栏" in t]
    assert left_idx and right_idx
    assert max(left_idx) < min(right_idx), "右栏内容出现在左栏之前"
    # 行缘证据在 crop 路径上命中
    assert any(ln.trailing_space for ln in lines), "左栏行尾空格证据丢失"
    assert any(ln.leading_space for ln in lines), "左栏行首空格证据丢失"
    # 端到端：两处折行空格均补回，且无假阳性插空
    paras = [b.text for b in analyze(pages).blocks if b.type == "paragraph"]
    assert any("中文与数字 2026 之间原有空格" in t for t in paras)
    assert any("以及中文与数字 2026 的又一次折行" in t for t in paras)
    assert not any("数字2026" in t for t in paras)


# ------------------------------------------------ DOCX 段落旗标
def test_docx_para_line_sets_flags_from_raw_text(tmp_path):
    """DOCX 行是显式段落（不跨段续接），旗标只服务标题折行合并路径；
    等宽行的前导空白已保留在 text 内，不再设 leading（否则重复补空）。"""
    import docx

    d = docx.Document()
    d.add_paragraph("普通段落行尾带空格 ")
    d.add_paragraph(" 普通段落行首带空格")
    p_mono = d.add_paragraph()
    run = p_mono.add_run("    code_line")
    run.font.name = "Consolas"
    p = tmp_path / "flags.docx"
    d.save(str(p))

    from doc2md.pipeline.s1_parse import read_docx
    lines = [ln for pg in read_docx(p) for ln in pg.lines]
    tail = next(ln for ln in lines if ln.text == "普通段落行尾带空格")
    lead = next(ln for ln in lines if ln.text == "普通段落行首带空格")
    mono = next(ln for ln in lines if "code_line" in ln.text)
    assert tail.trailing_space is True and tail.leading_space is False
    assert lead.leading_space is True and lead.trailing_space is False
    assert mono.leading_space is False, "等宽行前导空白已在 text 内，不得重复设旗"
    assert mono.text == "    code_line"


# ------------------------------------------------ 序列化纪律
def test_block_line_trail_not_serialized():
    """line_trail 是瞬态字段：不进 to_dict（SCHEMA_VERSION 不变），
    落库 / API 视图与既有产物逐字节一致。"""
    b = Block(id="b1", type="paragraph", text="数字 2026", line_trail=True)
    d = b.to_dict()
    assert "line_trail" not in d
    assert Block.from_dict(d).text == "数字 2026"
