"""CH-489 / 第十一轮全面审查的 A 类回归。

覆盖：
- 零页 PDF 不得被当成成功产物；
- DOCX 显式单行 / 单列表格保持表格结构；
- 显式 `limit=0` 不得绕过 500 行硬顶；
- `paths.tmp` 不得指向工作根或其祖先（否则 cron/serve 清扫会删掉工作根内容）；
- 删前快照必须 quick_check + 原子替换；
- watch 同步路径不得把附着到的既有作业提前标 done；
- 内容漂移任务必须指向实际产物，且同一批重复路径只读一次 IR、只落库一次；
- Markdown 链接 label/URL 转义，导出路径不接受路径段型 doc_id；
- LLM 候选 allowed_ops 必须与配置白名单求交。
"""
from __future__ import annotations

import copy
import importlib
import json
from pathlib import Path
from typing import ClassVar

import pytest

from doc2md.ir import IR, Block
from doc2md.scheduler import Job, Scheduler
from doc2md.storage.db import Store
from doc2md.storage.repo import save_ir

pytestmark = [pytest.mark.milestone("DoD")]


# --- 零页 PDF -----------------------------------------------------------------
@pytest.mark.gate
def test_zero_page_pdf_is_rejected_as_e4221(tmp_path, tmp_cfg):
    from pypdf import PdfWriter

    from doc2md.errors import E4221_GateReject
    from doc2md.pipeline.convert import convert

    src = tmp_path / "zero-pages.pdf"
    PdfWriter().write(str(src))
    with pytest.raises(E4221_GateReject) as ei:
        convert(src, cfg=tmp_cfg, out_dir=tmp_path / "out")
    detail = ei.value.detail
    assert any("no_pages" in r for r in detail.get("reasons", [])), detail


# --- DOCX 显式表格 ------------------------------------------------------------
@pytest.mark.gate
def test_docx_single_row_and_single_column_tables_stay_tables(tmp_path, tmp_cfg):
    import docx

    from doc2md.pipeline.convert import convert

    filler = ("这是一个用于通过文本层门禁的普通段落，包含足够多的中文字符以超过最小阈值。"
              * 2)
    d = docx.Document()
    d.add_paragraph(filler)
    one_row = d.add_table(rows=1, cols=2)
    one_row.rows[0].cells[0].text = "单行甲"
    one_row.rows[0].cells[1].text = "单行乙"
    one_col = d.add_table(rows=2, cols=1)
    one_col.rows[0].cells[0].text = "单列甲"
    one_col.rows[1].cells[0].text = "单列乙"
    d.add_paragraph(filler)
    src = tmp_path / "tables.docx"
    d.save(src)

    res = convert(src, cfg=tmp_cfg, out_dir=tmp_path / "out")
    tables = [b for b in res.ir.blocks if b.type == "table"]
    assert len(tables) == 2, [(b.type, b.text) for b in res.ir.blocks]
    assert "| 单行甲 | 单行乙 |" in tables[0].text
    assert "| 单列甲 |" in tables[1].text


# --- 分页硬顶 -----------------------------------------------------------------
@pytest.mark.gate
def test_explicit_zero_limit_is_capped_at_500(api_client, headers):
    client, tokens = api_client
    store = client.application.config["doc2md"]["store"]
    doc_id = "z" * 16
    ir = IR(doc_id=doc_id, source_key="z" * 24, source_name="cap.pdf", format="pdf",
            blocks=[Block(id=f"b{i}", type="paragraph", text=f"row-{i}")
                    for i in range(1, 601)])
    save_ir(store, ir)
    h = headers(tokens, "reader")
    body = client.get(f"/docs/{doc_id}?limit=0", headers=h).get_json()
    assert len(body["blocks"]) == 500 and body["blocks_total"] == 600
    assert body["blocks_limit"] == 500

    with store.tx() as tx:
        tx.executemany(
            "INSERT INTO edit_audit(doc_id, block_id, before, after, from_revision,"
            " to_revision, actor, ts, reason) VALUES (?,?,?,?,?,?,?,?,?)",
            [(doc_id, "b1", f"before-{i}", f"after-{i}", i + 1, i + 2, "t",
              "2026-01-01T00:00:00+00:00", "") for i in range(600)])
    audit = client.get(f"/docs/{doc_id}/audit?limit=0",
                       headers=headers(tokens, "editor")).get_json()
    assert len(audit["audit"]) == 500 and audit["total"] == 600

    with store.tx() as tx:
        tx.executemany(
            "INSERT INTO deletions(doc_id, source_key, actor, ts, kind, cascade_json)"
            " VALUES (?,?,?,?,?,?)",
            [(f"del-{i}", "sk", "t", "2026-01-01T00:00:00+00:00", "doc", "{}")
             for i in range(600)])
    dels = client.get("/audit/deletions?limit=0",
                      headers=headers(tokens, "admin")).get_json()
    assert len(dels["deletions"]) == 500 and dels["total"] == 600


# --- 配置路径安全 -------------------------------------------------------------
@pytest.mark.gate
def test_paths_tmp_must_not_swallow_the_work_root(tmp_cfg):
    from doc2md import config

    config.validate(tmp_cfg)                       # 默认 work/tmp 合法
    root = tmp_cfg["paths"]["root"]
    for bad in (root, str(Path(root).parent)):
        cfg = copy.deepcopy(tmp_cfg)
        cfg["paths"]["tmp"] = bad
        with pytest.raises(config.ConfigError):
            config.validate(cfg)


@pytest.mark.gate
def test_paths_out_is_runtime_immutable(tmp_cfg):
    """运行期改输出根会让旧产物在删除/导出两侧同时失联；与 root 同口径须重启。"""
    from doc2md import config

    candidate = copy.deepcopy(tmp_cfg)
    candidate["paths"]["out"] = str(Path(tmp_cfg["paths"]["root"]) / "other-out")
    with pytest.raises(config.ConfigError):
        config.check_runtime_immutable(candidate, tmp_cfg)
    config.check_runtime_immutable(copy.deepcopy(tmp_cfg), tmp_cfg)


# --- 删前快照 -----------------------------------------------------------------
@pytest.mark.gate
def test_predelete_snapshot_is_checked_and_atomic(tmp_path, tmp_cfg, monkeypatch):
    from doc2md.ops import backup

    root = Path(tmp_cfg["paths"]["root"])
    root.mkdir(parents=True, exist_ok=True)
    db = root / "state.sqlite"
    Store(db).close()
    backups = Path(tmp_cfg["ops"]["backup_local_dir"])
    plan = backup.BackupPlan(
        local_dir=backups, external_dir=None, rotation_days=14,
        work_root=root, work_dir_mode=str(tmp_cfg["ops"]["work_dir_mode"]),
        cfg=tmp_cfg)
    monkeypatch.setattr(backup, "_sqlite_quick_check", lambda _p: "error: corrupt")
    with pytest.raises(ValueError):
        plan.predelete_snapshot(db)
    assert not list(backups.glob("predelete-*")), list(backups.glob("*"))
    assert not list(backups.glob("*.part")), list(backups.glob("*"))


# --- watch 状态机 -------------------------------------------------------------
@pytest.mark.gate
@pytest.mark.parametrize(("attached", "expected_marks"), [(True, 0), (False, 1)])
def test_watch_marks_only_jobs_it_created(tmp_path, monkeypatch, attached,
                                          expected_marks):
    # 先导入 worker：cmd_watch 在函数体内 import worker；若首次导入发生在
    # repo.save_ir 被 patch 之后，worker 的模块级 save_ir 会永久绑定测试 lambda。
    from doc2md import cli
    from doc2md import worker as worker_mod  # noqa: F401
    from doc2md.storage import repo as repo_mod

    work_root = tmp_path / "work"
    work_root.mkdir()
    src = tmp_path / "a.pdf"
    src.write_bytes(b"%PDF-1.7\n")

    class FakeTracker:
        def __init__(self, stable_seconds: float = 0.0):
            self.stable_seconds = stable_seconds

        def poll_all(self, _d, _lim):
            return [str(src)]

    class FakeResult:
        ir = IR(doc_id="d" * 16, source_key="s" * 24, source_name="a.pdf",
                blocks=[Block(id="b1", text="正文")])
        meta: ClassVar[dict] = {}
        gate: ClassVar[dict] = {}
        doc_id = "d" * 16
        source_key = "s" * 24

    monkeypatch.setattr("doc2md.security.watcher.StabilityTracker", FakeTracker)
    convert_mod = importlib.import_module("doc2md.pipeline.convert")
    monkeypatch.setattr(convert_mod, "convert", lambda *_a, **_kw: FakeResult())
    monkeypatch.setattr(repo_mod, "save_ir", lambda *_a, **_kw: None)
    marks: list[tuple] = []
    monkeypatch.setattr(Scheduler, "submit", lambda _self, *_a, **_kw: Job(
        "job-watch", "d" * 16, "queued", attached=attached))
    monkeypatch.setattr(Scheduler, "mark",
                        lambda _self, *a, **kw: marks.append((a, kw)) or True)

    class Args:
        cmd, dir, once, interval, root = "watch", str(tmp_path), True, 0.0, str(work_root)
        config = None

    assert cli.cmd_watch(Args()) == cli.EXIT_OK
    assert len(marks) == expected_marks, marks


# --- 漂移任务 / 重复落库 -------------------------------------------------------
@pytest.mark.gate
def test_finish_ok_points_task_at_consumed_doc_on_drift(tmp_path):
    from doc2md.worker import _finish_ok

    store = Store(tmp_path / "state.sqlite")
    sched = Scheduler(store)
    j = sched.submit("e" * 16, "s" * 24, source_path="/tmp/x.pdf")
    job = {"job_id": j.job_id, "doc_id": "e" * 16, "source_key": "s" * 24,
           "path": "/tmp/x.pdf", "del_seq": 0, "force": False}
    ir = IR(doc_id="c" * 16, source_key="s" * 24, source_name="x.pdf",
            blocks=[Block(id="b1", text="正文")])
    res = {"doc_id": "c" * 16, "source_key": "s" * 24, "zip": "/tmp/x.zip",
           "blocks": 1, "code": ""}
    try:
        assert _finish_ok(store, sched, job, res, ir) is True
        row = sched.get(j.job_id)
        assert row["status"] == "done" and row["doc_id"] == "c" * 16, row
        detail = json.loads(row["detail_json"])
        assert detail["content_drift"]["enqueued"] == "e" * 16, detail
    finally:
        store.close()


@pytest.mark.gate
def test_drain_reads_and_saves_duplicate_path_once(tmp_path, tmp_cfg, monkeypatch):
    from doc2md import worker

    root = tmp_path / "work"
    (root / "logs").mkdir(parents=True)
    store = Store(root / "state.sqlite")
    sched = Scheduler(store)
    src = root / "x.pdf"
    src.write_bytes(b"%PDF-1.7\n")
    for i in range(3):
        sched.submit(f"dup{i}", "sk", source_path=str(src))

    ir = IR(doc_id="cons", source_key="sk", source_name="x.pdf",
            blocks=[Block(id="b1", text="正文")])
    monkeypatch.setattr(worker, "run_batch", lambda files, **_kw: [{
        "file": str(src), "status": "ok", "zip": "fake.zip", "doc_id": "cons",
        "blocks": 1, "code": "", "message": ""}])
    reads: list[str] = []
    saves: list[int] = []
    monkeypatch.setattr(worker, "read_ir", lambda z: reads.append(z) or ir)
    monkeypatch.setattr(worker, "save_ir", lambda *_a, **_kw: saves.append(1))
    try:
        rep = worker.drain_once(store, tmp_cfg, root, workers=1)
        assert rep.done == 3, rep.to_dict()
        assert len(reads) == 1, reads
        assert len(saves) == 1, saves
    finally:
        store.close()


# --- 渲染转义 / 导出路径 -------------------------------------------------------
@pytest.mark.gate
def test_markdown_link_label_and_url_are_safe():
    from doc2md.pipeline import s4_render

    ir = IR(doc_id="l" * 16, source_key="l" * 24, blocks=[
        Block(id="b1", text="[x]", links=[
            {"url": "https://example.com/a(b) c", "text": "[x]", "dest": ""}]),
        Block(id="b2", text="plain", links=[
            {"url": "https://example.com/x", "text": "plain", "dest": ""}]),
    ])
    md = s4_render.render(ir)
    assert r"[\[x\]](https://example.com/a%28b%29%20c)" in md, md
    assert "[plain](https://example.com/x)" in md, md


@pytest.mark.gate
def test_markdown_export_rebuilds_missing_artifact(api_client, headers, saved_doc,
                                                   tmp_cfg):
    client, tokens, doc_id = saved_doc
    zp = Path(tmp_cfg["paths"]["out"]) / f"{doc_id}.zip"
    assert zp.exists(), "saved_doc 应先落一份产物"
    zp.unlink()
    r = client.get(f"/docs/{doc_id}/markdown", headers=headers(tokens, "reader"))
    assert r.status_code == 200, r.get_data(as_text=True)[:300]
    assert zp.exists() and r.data


@pytest.mark.gate
def test_markdown_artifact_path_rejects_path_segments(tmp_path):
    from doc2md.api import _artifact_zip

    with pytest.raises(Exception):
        _artifact_zip(tmp_path, r"..\escape")
    with pytest.raises(Exception):
        _artifact_zip(tmp_path, "../escape")
    assert _artifact_zip(tmp_path, "safe-doc").name == "safe-doc.zip"


# --- LLM 白名单 ---------------------------------------------------------------
@pytest.mark.gate
def test_llm_candidate_allowed_ops_follow_configured_whitelist():
    from doc2md.llm import LlmClient

    blocks = [
        Block(id="b1", type="paragraph", text="这是一个没有句末标点的长段落内容"),
        Block(id="b2", type="paragraph", text="下一段正文"),
    ]
    client = LlmClient()
    one = client.payload_batches(blocks, whitelist=["merge_block"])[0]
    user = one["messages"][1]["content"]
    assert '"allowed_ops": ["merge_block"]' in user, user
    assert "白名单 op：[\"merge_block\"]" in user, user

    empty = client.payload_batches(blocks, whitelist=[])[0]
    empty_user = empty["messages"][1]["content"]
    assert '"allowed_ops": []' in empty_user, empty_user
    assert "白名单 op：[]" in empty_user, empty_user


# --- 旧库无 gate 也可从 DB 重建产物 -------------------------------------------
@pytest.mark.gate
def test_rebuild_artifact_without_gate_json(store, tmp_path):
    from doc2md.artifacts import rebuild_artifact

    doc_id = "g" * 16
    ir = IR(doc_id=doc_id, source_key="g" * 24, source_name="old.pdf",
            blocks=[Block(id="b1", text="旧库正文")])
    save_ir(store, ir)                       # 不传 gate → docs.gate_json = "{}"
    out = tmp_path / "out"
    p = rebuild_artifact(store, doc_id, {
        "paths": {"root": str(tmp_path), "out": str(out)}}, tmp_path)
    assert p is not None and p.exists(), "无 gate_json 的旧库也应能从 blocks 重建"
