"""N-2 批量压测探针（optimization-report §6 N-2）。

量化两个口径：
1. 单次 CLI 调用的固定开销（1 个小文件的耗时基线）；
2. 批量小文件下 --skip-existing 的收益随批量规模的摊薄
   （首转全量 vs 二转全 skipped，串行口径 workers=1 排除进程池干扰）。

只读探测：产物写入 tmp/n2_batch/，不触碰语料基线与 golden。
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import time
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
BATCH = ROOT / "tmp" / "n2_batch"
SRC = BATCH / "src"
OUT = BATCH / "out"
PY = sys.executable
CLI = [PY, "-m", "doc2md", "cli"]  # 占位，下方重拼


def make_corpus(n: int) -> Path:
    if SRC.exists():
        shutil.rmtree(SRC)
    SRC.mkdir(parents=True)
    for i in range(n):
        d = docx.Document()
        d.add_heading(f"压测样本 {i:04d}", level=1)
        d.add_paragraph(f"这是第 {i} 份压测样本的首段正文，用于 N-2 批量计时。")
        d.add_heading("小节标题", level=2)
        d.add_paragraph(
            f"样本 {i} 的第二段：包含若干文字以模拟真实小文档的体量，"
            "每份约两三段，便于测量固定开销占比。"
        )
        d.save(str(SRC / f"perf_{i:04d}.docx"))
    return SRC


def run_cli(args: list[str]) -> tuple[float, str]:
    t0 = time.perf_counter()
    proc = subprocess.run(
        [PY, "-c", "from doc2md.cli import main; main()", *args],
        capture_output=True, text=True, cwd=str(ROOT),
    )
    dt = time.perf_counter() - t0
    if proc.returncode != 0:
        raise SystemExit(f"CLI failed rc={proc.returncode}\n{proc.stderr[-2000:]}")
    return dt, proc.stdout


def parse_results(js: str) -> list[dict]:
    """CLI --json 为 JSONL（每行一条结果）。"""
    rows = []
    for line in js.splitlines():
        line = line.strip()
        if line.startswith("{"):
            rows.append(json.loads(line))
    return rows


def main() -> None:
    BATCH.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    make_corpus(200)

    results: dict[str, object] = {}

    # 1) 单文件固定开销（3 次取最小）
    single = SRC / "perf_0000.docx"
    singles = []
    out1 = BATCH / "out_single"
    if out1.exists():
        shutil.rmtree(out1)
    for _ in range(3):
        dt, js = run_cli(["convert", str(single), "--out", str(out1),
                          "--json", "--workers", "1"])
        singles.append(round(dt * 1000))
    results["single_file_ms_min"] = min(singles)
    results["single_file_ms_all"] = singles

    # 2) 200 文件首转（全量转换）
    dt1, js1 = run_cli(["convert", str(SRC), "--out", str(OUT),
                        "--json", "--workers", "1"])
    j1 = parse_results(js1)
    results["batch200_first_ms"] = round(dt1 * 1000)
    results["batch200_first_done"] = sum(
        1 for r in j1 if r.get("status") == "ok")

    # 3) 200 文件二转（--skip-existing，全部应 skipped）
    dt2, js2 = run_cli(["convert", str(SRC), "--out", str(OUT),
                        "--json", "--workers", "1", "--skip-existing"])
    j2 = parse_results(js2)
    skipped = sum(1 for r in j2 if r.get("status") == "skipped")
    results["batch200_skip_ms"] = round(dt2 * 1000)
    results["batch200_skipped"] = skipped
    if skipped != 200:
        raise SystemExit(f"expected 200 skipped, got {skipped}")

    # 4) 40 文件子集首转 vs 二转（规模对照）
    sub = BATCH / "src40"
    if sub.exists():
        shutil.rmtree(sub)
    sub.mkdir()
    for f in sorted(SRC.glob("*.docx"))[:40]:
        shutil.copy2(f, sub / f.name)
    out40 = BATCH / "out40"
    if out40.exists():
        shutil.rmtree(out40)
    dt3, _ = run_cli(["convert", str(sub), "--out", str(out40),
                      "--json", "--workers", "1"])
    dt4, js4 = run_cli(["convert", str(sub), "--out", str(out40),
                        "--json", "--workers", "1", "--skip-existing"])
    j4 = parse_results(js4)
    sk4 = sum(1 for r in j4 if r.get("status") == "skipped")
    results["batch40_first_ms"] = round(dt3 * 1000)
    results["batch40_skip_ms"] = round(dt4 * 1000)
    results["batch40_skipped"] = sk4
    if sk4 != 40:
        raise SystemExit(f"expected 40 skipped, got {sk4}")

    results["skip_ratio_40"] = round(
        results["batch40_skip_ms"] / results["batch40_first_ms"], 3)
    results["skip_ratio_200"] = round(
        results["batch200_skip_ms"] / results["batch200_first_ms"], 3)
    results["per_file_first_ms_200"] = round(
        results["batch200_first_ms"] / 200, 2)
    results["per_file_skip_ms_200"] = round(
        results["batch200_skip_ms"] / 200, 2)
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
