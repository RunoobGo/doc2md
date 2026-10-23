"""第二轮 QA 偏差回归（QA-D2M-2026-10-23-R2）。

D1–D13 逐项钉住修复后的行为。每项先于实现落地（TDD 红灯），
修复后本文件必须全绿；历史缺陷口径见 docs/qa/round2/comparison.md。
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import zipfile
from pathlib import Path

import pytest

from doc2md import PIPELINE_VERSION
from doc2md.cli import main as cli_main
from doc2md.edits import patch_block
from doc2md.ir import Block, IR
from doc2md.pipeline.convert import convert
from doc2md.pipeline.s3_structure import _gfm_table
from doc2md.pipeline.s4_render import assert_render_contract
from doc2md.storage.db import Store
from doc2md.storage.repo import save_ir

_FILLER = ("This paragraph exists to satisfy the S1.5 gate minimum characters per page "
           "threshold so that the conversion proceeds to the packaging stage. ")


def _make_pdf(path: Path, title: str, body: str) -> None:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    c.setFont("Helvetica", 12)
    c.drawString(72, 720, title)
    y = 690
    text = body + _FILLER * 3
    for i in range(0, len(text), 70):
        c.drawString(72, y, text[i:i + 70])
        y -= 18
        if y < 60:
            c.showPage()
            c.setFont("Helvetica", 12)
            y = 720
    c.save()


def _make_docx(path: Path, paragraphs: list[str]):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(path)
    return path


def _content_md(res) -> str:
    with zipfile.ZipFile(res.zip_path) as z:
        name = [n for n in z.namelist() if n.endswith("content.md")][0]
        return z.read(name).decode("utf-8")


# ---------------------------------------------------------------- D1 / D12
def test_d1_md_emit_failure_returns_error_status(tmp_path, capsys):
    src = tmp_path / "a.pdf"
    _make_pdf(src, "Failure", "Markdown emit failure path.")
    out = tmp_path / "out"
    out.mkdir()
    (out / "a.pdf.md").mkdir()          # 目标路径被目录占用 → os.replace 失败

    rc = cli_main(["convert", str(src), "--root", str(tmp_path / "root"),
                   "--out", str(out), "--format", "md", "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert rc == 4, f"md 直出失败必须是非零业务退出码，实际 {rc}"
    assert payload["status"] == "error"
    assert payload["code"] == "E5031"
    assert Path(payload["zip"]).is_file(), "失败时 zip 必须保留，便于调用方降级取件"


def test_d12_md_result_does_not_report_deleted_zip(tmp_path, capsys):
    src = tmp_path / "b.pdf"
    _make_pdf(src, "Markdown", "Zip removal reporting.")
    out = tmp_path / "out"
    rc = cli_main(["convert", str(src), "--root", str(tmp_path / "root"),
                   "--out", str(out), "--format", "md", "--json"])
    assert rc == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload.get("zip_removed") is True
    assert "zip" not in payload, "md 模式 zip 已删除，结果行不得再报路径"
    assert Path(payload["md"]).is_file()


# ---------------------------------------------------------------- D2 / D3
_LITERAL_PARAS = [
    "普通正文开头。",
    "# 这不是标题，只是井号开头的正文",
    "> 这不是引用，只是大于号开头的正文",
    "[点击这里](https://example.com) 在源文档中是纯文本",
    "字面星号：**这不是加粗**",
    "1. 这不是列表，而是以编号开头的正文段",
]


def test_d2_markdown_literals_stay_plain_text(tmp_path):
    src = _make_docx(tmp_path / "literals.docx", _LITERAL_PARAS)
    res = convert(src, out_dir=tmp_path / "out")
    types = [b.type for b in res.ir.blocks]
    assert types == ["paragraph"] * len(_LITERAL_PARAS), \
        f"Markdown 字面量不得新增结构：{[ (b.type, b.text) for b in res.ir.blocks ]}"

    md = _content_md(res)
    assert r"\# 这不是标题" in md
    assert r"\> 这不是引用" in md
    assert r"\[点击这里\](https://example.com)" in md
    assert r"\*\*这不是加粗\*\*" in md
    assert r"1\. 这不是列表" in md


def test_d3_raw_html_and_javascript_are_escaped(tmp_path):
    paras = [
        "<script>alert(1)</script>",
        "<img src=x onerror=alert(2)>",
        "[点我](javascript:alert(3))",
        "普通正文。" + _FILLER,
    ]
    src = _make_docx(tmp_path / "html.docx", paras)
    res = convert(src, out_dir=tmp_path / "out")
    md = _content_md(res)

    assert re.search(r"(?<!\\)<script", md) is None, md
    assert re.search(r"(?<!\\)<img", md) is None, md
    assert r"\<script" in md and r"\<img" in md
    assert re.search(r"(?<!\\)\]\(javascript:", md) is None, md
    assert r"\[点我\](javascript:alert(3))" in md
    assert_render_contract(md)              # 产物自身必须过渲染红线

    for bad in ("<script>alert(1)</script>", "<img src=x onerror=alert(2)>",
                "[x](javascript:alert(1))"):
        with pytest.raises(AssertionError):
            assert_render_contract(bad)


# ---------------------------------------------------------------- D4 / D5
def test_d4_table_cell_escapes_backslash_before_pipe():
    md = _gfm_table([["甲", "乙", "丙"],
                     [r"反斜杠\|组合", "普通", "末列"]])
    assert r"反斜杠\\\|组合" in md, md
    assert "| 末列 |" in md


def test_d4_docx_cell_survives_gfm_round_trip(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("单元格转义测试。")
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(("列A", "列B", "列C")):
        t.cell(0, i).text = v
    t.cell(1, 0).text = r"反斜杠\|组合"
    t.cell(1, 1).text = "管道|字符"
    t.cell(1, 2).text = "字面<br>标记"
    src = tmp_path / "cell.docx"
    d.save(src)

    res = convert(src, out_dir=tmp_path / "out")
    table = [b for b in res.ir.blocks if b.type == "table"][0]
    assert r"反斜杠\\\|组合" in table.text
    assert r"管道\|字符" in table.text
    assert "字面&lt;br&gt;标记" in table.text, "源文本里的 <br> 是字面量，不得当结构换行"


def test_d5_nested_table_is_flat_counted_and_warned(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("嵌套表格测试。")
    outer = d.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "外层左"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "内层A"
    inner.cell(0, 1).text = "内层B"
    src = tmp_path / "nested.docx"
    d.save(src)

    res = convert(src, out_dir=tmp_path / "out")
    table = [b for b in res.ir.blocks if b.type == "table"][0]
    assert table.flags.get("flattened_cells", 0) >= 1
    assert "W1045" in (res.meta.get("warnings") or []), \
        f"嵌套表拍平必须留 W1045：{res.meta.get('warnings')}"


# ---------------------------------------------------------------- D6
def test_d6_page_break_inside_paragraph_splits_at_position(tmp_path):
    import docx
    from docx.enum.text import WD_BREAK
    d = docx.Document()
    d.add_paragraph("分页前的独立段落。")
    p = d.add_paragraph()
    p.add_run("同一段内分页前文字。")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("同一段内分页后文字。")
    d.add_paragraph("分页后的下一段正文。")
    src = tmp_path / "pagebreak.docx"
    d.save(src)

    res = convert(src, out_dir=tmp_path / "out")
    seq = [(b.type, b.text) for b in res.ir.blocks]
    assert ("paragraph", "同一段内分页前文字。") in seq
    assert ("paragraph", "同一段内分页后文字。") in seq
    i = seq.index(("paragraph", "同一段内分页前文字。"))
    j = seq.index(("paragraph", "同一段内分页后文字。"))
    assert j == i + 2 and seq[i + 1][0] == "page_break", seq


# ---------------------------------------------------------------- D7
def test_d7_midpage_large_callout_is_not_heading():
    from doc2md.pipeline.s3_structure import _heading_level

    callout = "以下一行是 14pt 的提示框文字，没有句末标点"
    assert _heading_level(callout, 14.0, 10.5, False,
                          top=121.8, page_height=841.89) is None
    # 正向：页顶短行 / 粗体仍可判标题
    assert _heading_level("代码说明", 14.0, 10.0, False,
                          top=46.5, page_height=841.89) == 2
    assert _heading_level("重点提示", 14.0, 10.5, True,
                          top=300.0, page_height=841.89) == 2


def test_d7_pdf_callout_stays_paragraph(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    cjk = "STSong-Light"
    path = tmp_path / "callout.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    body = ["这是正文段落，用于建立 10.5 磅正文字号基准并满足门禁字符数要求。"] * 12

    c.setFont(cjk, 10.5)
    c.drawString(72, 760, "第一章 系统概述")
    y = 730
    for t in body:
        c.drawString(72, y, t)
        y -= 18
    c.setFont(cjk, 14)
    c.drawString(72, 610, "以下一行是 14pt 的提示框文字，没有句末标点")
    c.setFont(cjk, 10.5)
    for t in body:
        c.drawString(72, y, t)
        y -= 18
    c.showPage()

    c.setFont(cjk, 14)
    c.drawString(72, 760, "结语")
    c.setFont(cjk, 10.5)
    y = 730
    for t in body:
        c.drawString(72, y, t)
        y -= 18
    c.save()

    res = convert(path, out_dir=tmp_path / "out")
    headings = [b.text for b in res.ir.blocks if b.type == "heading"]
    assert "以下一行是 14pt 的提示框文字，没有句末标点" not in headings, headings
    assert "结语" in headings, headings
    assert any(b.type == "paragraph" and "提示框文字" in (b.text or "")
               for b in res.ir.blocks)


# ---------------------------------------------------------------- D8
def test_d8_same_text_patch_is_noop(store):
    ir = IR(doc_id="d1", source_key="s1", source_name="a.pdf", format="pdf", pages=1)
    ir.blocks = [Block(id="b1", type="paragraph", text="original text")]
    save_ir(store, ir)

    out = patch_block(store, "d1", "b1",
                      {"text": "original text", "base_revision": 1}, "tester")
    assert out["unchanged"] is True
    assert out["revision"] == 1
    assert store.q1("SELECT COUNT(*) c FROM edit_audit")["c"] == 0
    row = store.q1("SELECT revision, text FROM blocks WHERE doc_id='d1' AND block_id='b1'")
    assert (row["revision"], row["text"]) == (1, "original text")

    out2 = patch_block(store, "d1", "b1",
                       {"text": "changed text", "base_revision": 1}, "tester")
    assert out2["revision"] == 2 and out2.get("unchanged") is None
    assert store.q1("SELECT COUNT(*) c FROM edit_audit")["c"] == 1


# ---------------------------------------------------------------- D9
def test_d9_explicit_zero_limit_is_empty_everywhere(api_client, headers):
    client, tokens = api_client
    st = client.application.config["doc2md"]["store"]
    ir = IR(doc_id="z" * 16, source_key="z" * 24, source_name="zero.pdf",
            format="pdf", pages=1,
            blocks=[Block(id=f"b{i}", type="paragraph", text="alpha zero")
                    for i in range(3)])
    save_ir(st, ir)

    h = headers(tokens, "reader")
    body = client.get(f"/docs/{ir.doc_id}?limit=0", headers=h).get_json()
    assert body["blocks"] == [] and body["blocks_limit"] == 0

    with st.tx() as tx:
        tx.execute(
            "INSERT INTO edit_audit(doc_id, block_id, before, after, from_revision,"
            " to_revision, actor, ts, reason) VALUES (?,?,?,?,?,?,?,?,?)",
            (ir.doc_id, "b1", "a", "b", 1, 2, "t", "2026-01-01T00:00:00+00:00", ""))
    audit = client.get(f"/docs/{ir.doc_id}/audit?limit=0",
                       headers=headers(tokens, "editor")).get_json()
    assert audit["audit"] == [] and audit["limit"] == 0

    with st.tx() as tx:
        tx.execute(
            "INSERT INTO deletions(doc_id, source_key, actor, ts, kind, cascade_json)"
            " VALUES (?,?,?,?,?,?)",
            ("del-1", "sk", "t", "2026-01-01T00:00:00+00:00", "doc", "{}"))
    dels = client.get("/audit/deletions?limit=0",
                      headers=headers(tokens, "admin")).get_json()
    assert dels["deletions"] == [] and dels["limit"] == 0

    jobs = client.get("/jobs?limit=0", headers=h).get_json()
    assert jobs["tasks"] == []
    hits = client.get("/search?q=alpha&limit=0", headers=h).get_json()
    assert hits["hits"] == []


def test_d9_cli_docs_list_zero_limit_is_empty(tmp_path, capsys):
    root = tmp_path / "root"
    root.mkdir()
    st = Store(root / "state.sqlite")
    try:
        save_ir(st, IR(doc_id="a" * 16, source_key="s" * 24, source_name="one.pdf",
                       format="pdf", pages=1,
                       blocks=[Block(id="b1", type="paragraph", text="正文")]))
    finally:
        st.close()
    rc = cli_main(["docs", "list", "--root", str(root), "--limit", "0"])
    assert rc == 0
    out = json.loads(capsys.readouterr().out)
    assert out["docs"] == []
    assert out["limit"] == 0
    assert out["has_more"] is False


# ---------------------------------------------------------------- D10
def test_d10_malformed_override_maps_to_config_error(tmp_path, capsys):
    root = tmp_path / "work"
    root.mkdir()
    (root / "settings_override.json").write_text("{not valid json", encoding="utf-8")
    rc = cli_main(["serve", "--root", str(root)])
    assert rc == 2, "配置错误必须按用法错误退出（2），不得落 E5000/5"
    err = json.loads(capsys.readouterr().err)
    assert err["code"] == "E4001"
    assert "settings_override.json" in err["message"]
    assert "line" in err["message"].lower()


# ---------------------------------------------------------------- D11
def test_d11_readme_version_matches_pipeline():
    text = Path("README.md").read_text(encoding="utf-8")
    assert f"v{PIPELINE_VERSION}" in text.splitlines()[2]
    assert "v2.12.1" not in text


# ---------------------------------------------------------------- D13
def test_d13_worker_event_path_is_locatable(store):
    from doc2md import worker
    path = ("/Volumes/fx/Object/Doc2md/work/inbox/2026/10/23/"
            "pdf_01_simple_1p.pdf")
    worker._event(store, {"job_id": "j1", "doc_id": "d1", "path": path},
                  {"blocks": 1}, "job_done", None)
    row = store.q1("SELECT payload_json FROM events WHERE code='job_done'")
    payload = json.loads(row["payload_json"])
    # 实现按 `os.path.abspath` 取平台绝对路径——Windows runner 上 /Volumes/...
    # 会落成 D:\Volumes\...，断言必须同口径计算，否则测的是宿主而不是契约。
    digest = hashlib.sha256(
        os.path.abspath(path).encode("utf-8")).hexdigest()[:8]
    assert payload["file"] == f"pdf_01_simple_1p.pdf#{digest}", payload["file"]
    assert len(payload["file"]) <= 64, "必须短于 scrub_record 的指纹化阈值"
