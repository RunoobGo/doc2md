"""CH-485 / 管线结构：10.5pt、表格、编号、文本框、marks、门禁、链接、分片。"""
from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

import pytest

from doc2md.pipeline.s1_parse import Line, RawPage, read_docx, read_pdf
from doc2md.pipeline.s3_structure import analyze

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = f'xmlns:w="{W}"'
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _lines(path: Path) -> list[Line]:
    return [ln for pg in read_docx(path) for ln in pg.lines]


def _inject_numbering(path: Path, paragraphs: list[tuple[int, str]],
                      numbering_xml: str | None = None) -> None:
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    body = "".join(
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{lvl}"/>'
        f'<w:numId w:val="1"/></w:numPr></w:pPr>'
        f'<w:r><w:t>{text}</w:t></w:r></w:p>'
        for lvl, text in paragraphs)
    items["word/document.xml"] = items["word/document.xml"].decode().replace(
        "</w:body>", body + "</w:body>").encode()
    items["word/numbering.xml"] = (numbering_xml or _NUMBERING).encode()
    rels = items["word/_rels/document.xml.rels"].decode()
    if "numbering.xml" not in rels:
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rIdNum1" Type="{R_NS}/numbering" Target="numbering.xml"/>'
            "</Relationships>")
        items["word/_rels/document.xml.rels"] = rels.encode()
    ct = items["[Content_Types].xml"].decode()
    if "numbering.xml" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/word/numbering.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
            "</Types>")
        items["[Content_Types].xml"] = ct.encode()
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, v in items.items():
            z.writestr(n, v)
    tmp.replace(path)


_DECIMAL_NUMBERING = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="{W}">
 <w:abstractNum w:abstractNumId="0">
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
 </w:abstractNum>
 <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
</w:numbering>'''


_NUMBERING = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="{W}">
 <w:abstractNum w:abstractNumId="0">
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/></w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
 </w:abstractNum>
 <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
</w:numbering>'''


@pytest.mark.gate
def test_105pt_numbered_line_is_list_not_heading():
    from reportlab.pdfgen import canvas
    path = Path(tempfile.mkdtemp()) / "p105.pdf"
    c = canvas.Canvas(str(path))
    y = 760
    for text in ("This is ordinary body text used to establish the body size." * 2,
                 "1. First item", "This line continues the ordinary body text.",
                 "More ordinary body text at the same size." * 2):
        c.setFont("Helvetica", 10.5)
        c.drawString(72, y, text)
        y -= 18
    c.save()
    pages = read_pdf(path, chunk_pages=0)
    ir = analyze(pages)
    assert any(b.type == "list_item" and "First item" in b.text for b in ir.blocks)
    cjk = analyze([RawPage(index=0, lines=[
        Line("第 1 章的内容需要结合上下文理解", size=10.5),
        Line("正文。", size=10.5),
    ])])
    assert cjk.blocks[0].type == "paragraph"


@pytest.mark.gate
def test_docx_numbered_table_first_column_stays_table(tmp_path):
    import docx
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("1.1 范围", "内容甲"),
                                ("1.2 目标", "内容乙"),
                                ("普通", "内容丙")]):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b
    path = tmp_path / "table.docx"
    d.save(str(path))
    ir = analyze(read_docx(path))
    assert any(b.type == "table" for b in ir.blocks)
    assert not any(b.type == "heading" for b in ir.blocks)


@pytest.mark.gate
def test_docx_bullet_and_multilevel_numbering(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("正文")
    path = tmp_path / "num.docx"
    d.save(str(path))
    _inject_numbering(path, [(0, "项目甲"), (0, "项目乙"), (1, "子项")])
    lines = _lines(path)
    numbered = [ln for ln in lines if ln.list_marker]
    assert [(ln.list_level, ln.list_marker) for ln in numbered] == [
        (0, "•"), (0, "•"), (1, "2.1.")], numbered
    from doc2md.pipeline.s4_render import render
    ir = analyze([RawPage(index=0, lines=lines)])
    md = render(ir)
    assert "- 项目甲" in md and "- 项目乙" in md
    assert "2.1. 子项" in md and "-   2.1." not in md


@pytest.mark.gate
def test_docx_heading_keeps_auto_number(tmp_path):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    p = d.add_heading("方法", level=1)
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    path = tmp_path / "heading_num.docx"
    d.save(str(path))
    _inject_numbering(path, [], numbering_xml=_DECIMAL_NUMBERING)
    ir = analyze(read_docx(path))
    heads = [b for b in ir.blocks if b.type == "heading"]
    assert heads and heads[0].text.startswith("1. "), heads


@pytest.mark.gate
def test_docx_textbox_content_is_preserved(tmp_path):
    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("前")
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')
    p._p.append(parse_xml(
        f'<w:r {ns}><w:drawing><wp:inline><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>框内文字</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>'
        '</wp:inline></w:drawing></w:r>'))
    p.add_run("后")
    path = tmp_path / "tb.docx"
    d.save(str(path))
    texts = [ln.text for ln in _lines(path)]
    assert "前后" in texts and "框内文字" in texts


def test_marks_are_and_across_continuation():
    pages = [RawPage(index=0, lines=[
        Line("无句末的粗体段落", size=10.0, bold=True, uniform_bold=True),
        Line("", size=10.0),
        Line("，普通续行", size=10.0, uniform_bold=False),
    ])]
    ir = analyze(pages)
    assert not ir.blocks[-1].marks.get("bold")


@pytest.mark.gate
def test_numeric_table_pages_pass_auto_gate():
    from doc2md.pipeline.s1_5_gate import check
    pages = [RawPage(index=i, lines=[
        Line("2024 100 250 390 120 450", size=10.0) for _ in range(20)])
        for i in range(2)]
    assert check(pages, min_chars=10, mode="auto").passed


@pytest.mark.gate
def test_pdf_multiline_link_not_expanded_to_whole_block():
    from reportlab.pdfgen import canvas
    path = Path(tempfile.mkdtemp()) / "link.pdf"
    c = canvas.Canvas(str(path))
    c.setFont("Helvetica", 10.5)
    c.drawString(72, 720, "See URL for the details of this agreement")
    c.linkURL("https://example.com/x", (72, 715, 300, 732), relative=0)
    c.drawString(72, 704, "and the second line continues the paragraph")
    c.save()
    ir = analyze(read_pdf(path, chunk_pages=0))
    # 链接行独立成块：URL 仍可点，但只覆盖首行，绝不包住续行。
    assert ir.blocks[0].text.startswith("See URL")
    assert ir.blocks[0].links and ir.blocks[0].links[0]["url"] == "https://example.com/x"
    assert all(not b.links for b in ir.blocks[1:])


def test_cross_shard_list_items_are_not_merged():
    from doc2md.ir import IR, Block
    from doc2md.pipeline.sharding import stitch_shards
    a = IR(doc_id="a", source_key="a", blocks=[Block(
        id="b1", type="list_item", text="1. 第一项",
        font={"family": "serif", "size": 10.0}, page=0)],
        extra_meta={"shard": {"start": 0, "end": 1}})
    b = IR(doc_id="a", source_key="a", blocks=[Block(
        id="b1", type="list_item", text="2. 第二项",
        font={"family": "serif", "size": 10.0}, page=1)],
        extra_meta={"shard": {"start": 1, "end": 2}})
    out, stats = stitch_shards([a, b])
    assert stats["stitched"] == 0 and len(out.blocks) == 2


@pytest.mark.gate
def test_batch_error_message_does_not_echo_exception_text(monkeypatch):
    import importlib

    import doc2md.pipeline.batch as batch
    conv = importlib.import_module("doc2md.pipeline.convert")

    def boom(*a, **kw):
        raise RuntimeError("SECRET-BODY")

    monkeypatch.setattr(conv, "convert", boom)
    r = batch._convert_one("/tmp/x.pdf", None, None, None)
    assert "SECRET-BODY" not in r.get("message", "")
    assert r["code"] == "E5000"
