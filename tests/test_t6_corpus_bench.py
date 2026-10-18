"""T6：大语料实测脚本的仪器自检 + CH-457 打开边界回归。

两部分：
1. CH-457 回归——魔术码合法但结构损坏的源文件（截断 PDF / 坏包 DOCX）在 S1
   打开边界收敛为 415/E4151。T6 大语料实测暴露：原实现未映射，pdfplumber 的
   PdfminerException 冒泡成 E5000 内部错误，「输入损坏」被误报成「工具故障」
   （CLI 退出码 5、API 500、任务终态 failed 且错误码不可判定）。
   覆盖 read_pdf / read_docx / convert / worker 落库全链路，以及好文件不误伤。
2. 基准仪器自检——语料字节级确定（同种子两次生成逐字节一致）、坏文件字节唯一
   （同字节会被 uq_tasks_doc_active 附着去重，失败率表口径失真）、三张表渲染、
   多后端进程树采样原语（proc / ps / psutil / none，T3 跨机口径）、run_drain
   真 CLI 端到端。测量工具自身出错比测出坏数字更危险。
"""
from __future__ import annotations

import json
import os
import sqlite3
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "ops" / "bench"))

import t6_large_corpus as bench  # noqa: E402

from doc2md import worker  # noqa: E402
from doc2md.config import load  # noqa: E402
from doc2md.errors import E4151_UnsupportedType  # noqa: E402
from doc2md.pipeline.convert import convert  # noqa: E402
from doc2md.pipeline.s1_parse import _open_source, read_docx, read_pdf, sniff  # noqa: E402
from doc2md.scheduler import Scheduler  # noqa: E402
from doc2md.storage.db import Store  # noqa: E402


# --- 夹具 ----------------------------------------------------------------------
@pytest.fixture
def truncated_pdf(tmp_path) -> Path:
    """魔术码合法、无 xref/trailer 的截断 PDF（sniff 通过，S1 打开失败）。"""
    p = tmp_path / "truncated.pdf"
    p.write_bytes(b"%PDF-1.4\n" + b"x" * 172)
    return p


@pytest.fixture
def bad_docx(tmp_path) -> Path:
    """zip 含 word/document.xml 但内容非 XML（sniff 通过，python-docx 打开失败）。"""
    p = tmp_path / "bad.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("word/document.xml", "<not valid xml <<<")
    return p


@pytest.fixture
def good_pdf(tmp_path) -> Path:
    from reportlab.pdfgen import canvas as rl_canvas
    p = tmp_path / "good.pdf"
    c = rl_canvas.Canvas(str(p), pagesize=(595, 842), invariant=1)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 800, "Bench Document good")
    c.setFont("Helvetica", 10.5)
    # 正文要过 S1.5 门禁（min_chars=120）：中英混排、足量字符
    body = ("转换 管线 离线 分片 缝合 门禁 备份 加密 队列 渲染 指纹 消毒 "
            "doc2md pipeline offline shard stitch gate backup encrypt. ") * 2
    y = 760
    for seg in range(0, len(body), 60):
        c.drawString(72, y, body[seg:seg + 60])
        y -= 16
    c.showPage()
    c.save()
    return p


@pytest.fixture
def good_docx(tmp_path) -> Path:
    from docx import Document
    p = tmp_path / "good.docx"
    d = Document()
    d.add_heading("Good Report", 0)
    d.add_paragraph("转换 管线 离线 分片 缝合 门禁。" * 4)
    bench._save_docx_det(d, p)
    return p


@pytest.fixture
def notpdf(tmp_path) -> Path:
    p = tmp_path / "notpdf.pdf"
    p.write_text("这不是 PDF，只是文本。\n" * 4, encoding="utf-8")
    return p


# --- CH-457：打开边界回归 --------------------------------------------------------
class TestOpenBoundaryE4151:
    """结构损坏的源文件必须在打开边界收敛为 415/E4151（方案 附录 D 冻结语义：
    magic 判定不支持的类型（含不可读）；打开源文件 OSError 同归此码）。"""

    def test_truncated_pdf_maps_to_e4151(self, truncated_pdf):
        assert sniff(truncated_pdf) == "pdf"          # 魔术码合法，过 sniff
        with pytest.raises(E4151_UnsupportedType) as ei:
            read_pdf(truncated_pdf)
        exc = ei.value
        assert exc.code == "E4151"
        assert exc.http == 415
        # detail 只含文件名与异常类名：异常消息可能夹带文档内容（NFR-20 不带入）
        assert set(exc.detail) == {"file", "exc"}
        assert exc.detail["file"] == truncated_pdf.name
        assert isinstance(exc.detail["exc"], str)
        assert "Unexpected EOF" not in json.dumps(exc.detail, ensure_ascii=False)

    def test_truncated_pdf_convert_surface_e4151(self, truncated_pdf, tmp_path):
        # convert 面同样收敛：不再是 E5000 内部错误
        with pytest.raises(E4151_UnsupportedType):
            convert(truncated_pdf, cfg=load(), out_dir=tmp_path / "out")

    def test_unopenable_docx_maps_to_e4151(self, bad_docx):
        assert sniff(bad_docx) == "docx"
        with pytest.raises(E4151_UnsupportedType) as ei:
            read_docx(bad_docx)
        assert ei.value.detail["file"] == bad_docx.name

    def test_good_files_not_misjudged(self, good_pdf, good_docx, tmp_path):
        pages = read_pdf(good_pdf)
        assert len(pages) == 1 and pages[0].chars > 0
        assert len(read_docx(good_docx)) >= 1
        res = convert(good_pdf, cfg=load(), out_dir=tmp_path / "out")
        assert res.ir.blocks and res.zip_path.exists()

    def test_open_source_passes_doc2mderror_through(self):
        original = E4151_UnsupportedType("already classified")
        def raiser():
            raise original
        with pytest.raises(E4151_UnsupportedType) as ei:
            _open_source(raiser, "x.pdf", "pdf")
        assert ei.value is original              # 不重包装，保持原异常对象

    def test_worker_marks_truncated_task_failed_e4151(self, truncated_pdf,
                                                      notpdf, tmp_path):
        """端到端落库：截断 PDF 入队成功（sniff 过）→ 消费失败 → failed/E4151；
        非文档文件入队即拒（E4151）且 queued 计数零增长（方案 12 不变式）。"""
        root = tmp_path / "work"
        root.mkdir()
        db = root / "state.sqlite"
        store = Store(db)
        sched = Scheduler(store)
        cfg = load()
        try:
            it = worker.enqueue(sched, truncated_pdf)
            assert it["status"] == "queued"
            with pytest.raises(E4151_UnsupportedType):
                worker.enqueue(sched, notpdf)
            rep = worker.drain_once(store, cfg, root, workers=1)
            assert (rep.drained, rep.failed) == (1, 1)
        finally:
            store.close()
        con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
        try:
            rows = con.execute("SELECT status, error_code FROM tasks").fetchall()
        finally:
            con.close()
        assert rows == [("failed", "E4151")]


# --- 基准仪器自检 ----------------------------------------------------------------
class TestCorpusGeneration:
    def test_byte_deterministic_same_seed(self, tmp_path):
        """同种子两次生成逐字节一致（12.9「固定语料」）。reportlab 默认嵌
        CreationDate、python-docx zip 条目带当前时间，均已固定。"""
        a, b = tmp_path / "a", tmp_path / "b"
        bench.generate_corpus(a, n_pdf=3, n_docx=2, n_bad=6, seed=20261014)
        bench.generate_corpus(b, n_pdf=3, n_docx=2, n_bad=6, seed=20261014)
        files = sorted(p for p in a.rglob("*") if p.is_file())
        assert len(files) == 11
        for p in files:
            assert p.read_bytes() == (b / p.relative_to(a)).read_bytes(), p

    def test_seed_changes_content(self, tmp_path):
        a, b = tmp_path / "a", tmp_path / "b"
        bench.generate_corpus(a, n_pdf=3, n_docx=2, n_bad=6, seed=1)
        bench.generate_corpus(b, n_pdf=3, n_docx=2, n_bad=6, seed=2)
        pa = (a / "pdf" / "small-00000.pdf").read_bytes()
        pb = (b / "pdf" / "small-00000.pdf").read_bytes()
        assert pa != pb

    def test_bad_files_byte_unique(self, tmp_path):
        """同字节文件会被 uq_tasks_doc_active 附着去重（约定 14）——
        失败率表的「文件数 ↔ 终态」1:1 口径要求坏文件字节唯一。空文件在 sniff
        即拒、根本到不了去重，故只要求截断/错码两类唯一。"""
        root = tmp_path / "corpus"
        bench.generate_corpus(root, n_pdf=1, n_docx=1, n_bad=9, seed=20261014)
        bad = sorted((root / "bad").iterdir())
        assert len(bad) == 9
        for prefix in ("truncated", "notpdf"):
            group = [p for p in bad if p.name.startswith(prefix)]
            assert len({p.read_bytes() for p in group}) == len(group) == 3
        # 截断 PDF 魔术码合法（过 sniff，故障点在 S1 打开）；另两类 sniff 即拒
        trunc = [p for p in bad if p.name.startswith("truncated")]
        for p in trunc:
            assert sniff(p) == "pdf"
        for p in bad:
            if p.name.startswith(("notpdf", "empty")):
                with pytest.raises(E4151_UnsupportedType):
                    sniff(p)


class TestTables:
    def _result(self, workers: int) -> dict:
        return {
            "workers": workers, "exit": 5, "wall_s": 120.0,
            "final": {"done": 51, "failed": 1},
            "rejected_codes": {"E4151": 1},
            "enqueue_rejected": {"E4151": 2},
            "sampler": "proc", "rss_available": True,
            "peak_tree_mb": 148.8, "peak_single_mb": 62.4,
            "backlog_peak": 54, "half_drain_s": 61.0,
            "queue_series": [], "rss_series": [],
        }

    def test_three_tables_rendered(self):
        info = {"pdf": 48, "docx": 3, "bad": 3, "total": 54, "seed": 20261014}
        out = bench.tables([self._result(1), self._result(4)], info)
        for header in ("T6-1 队列积压", "T6-2 驻留", "T6-3 失败率"):
            assert header in out
        assert "workers=1" in out and "workers=4" in out
        assert "| 排空耗时（s） |" in out
        assert "| 进程树峰值 RSS（MB） |" in out
        assert "| 采样后端 | proc | proc |" in out
        assert "| 非 done 占比 |" in out
        # 3/54 非 done → 5.56%
        assert "5.56%" in out
        # 错误码分解合并终态与入队即拒两侧
        assert "| ├─ E4151 | 3 | 3 |" in out

    def test_rss_na_when_sampler_unavailable(self):
        """平台无进程树采样能力时输出 n/a，不得输出 0（0 会被读成「不驻留」）。"""
        r = self._result(1)
        r["sampler"], r["rss_available"] = "none", False
        r["peak_tree_mb"] = r["peak_single_mb"] = None
        info = {"pdf": 48, "docx": 3, "bad": 3, "total": 54, "seed": 20261014}
        out = bench.tables([r], info)
        assert "| 采样后端 | none |" in out
        assert "| 进程树峰值 RSS（MB） | n/a" in out
        assert "| 最大单进程 RSS（MB） | n/a" in out
        # 吞吐与失败账目不受采样能力影响
        assert "| 排空耗时（s） | 120.0 |" in out
        assert "5.56%" in out


class TestSamplerPrimitives:
    def test_default_backend_samples_tree(self):
        """CH-469：默认后端断言改为平台无关——写死 "proc" 等于把「本机是 Linux」
        编进断言，macOS 上必红（实测）。契约是：POSIX 本机必然有可用后端，
        且能采到自身 pid 与正的 RSS。
        """
        if os.name == "nt" and bench._probe_backend() == "none":
            pytest.skip("Windows without psutil: RSS sampling is intentionally n/a")
        pids, rss, backend = bench.sample_tree(os.getpid())
        assert backend in ("proc", "ps", "psutil")       # POSIX 探测结果因宿主而异
        assert os.getpid() in pids
        assert rss.get(os.getpid(), 0) > 0

    def test_proc_branch_reads_a_fake_proc_tree(self, tmp_path):
        """proc 分支在任意宿主上确定测试：CH-469 把 `/proc` 根做成可注入参数。

        原来的 `pytest.skip("host has no /proc")` 让 macOS 上这条分支零覆盖，
        而它恰恰是 Linux 长跑用的采样路径——跳过等于把「本机不是 Linux」当成
        「不必验证」。现在用假 proc 树断言：后代遍历（含孙进程）、VmRSS 解析、
        非数字目录忽略、status 缺失的进程跳过而不是整棵崩。
        """
        root, child, grand = os.getpid(), os.getpid() + 1, os.getpid() + 2
        for pid, ppid, rss in ((root, 1, None), (child, root, 2048), (grand, child, 4096)):
            d = tmp_path / str(pid)
            d.mkdir()
            lines = [f"Name:\tp{pid}", f"PPid:\t{ppid}"]
            if rss is not None:
                lines.append(f"VmRSS:\t{rss} kB")
            (d / "status").write_text("\n".join(lines) + "\n", encoding="ascii")
        (tmp_path / "self").mkdir()                    # 非数字条目：忽略
        vanished = tmp_path / str(grand + 1)           # 采样瞬间退出的进程：跳过
        vanished.mkdir()                               # 有 pid 目录、无 status 可读

        pids, rss = bench._tree_proc(root, proc_root=str(tmp_path))
        assert sorted(pids) == sorted([root, child, grand])
        assert rss == {child: 2048, grand: 4096}       # root 无 VmRSS 行 → 不入表

    def test_proc_backend_only_when_proc_exists(self):
        """真实 /proc 存在时（Linux CI）确认探测与采样口径一致；非 Linux 宿主
        由 `test_proc_branch_reads_a_fake_proc_tree` 覆盖同一分支，不产生跳过。
        """
        if not os.path.isdir("/proc"):
            if os.name == "nt":
                pytest.skip("Windows has no /proc and no ps backend")
            assert bench._probe_backend() == "ps"       # macOS：POSIX 无 proc → ps
        else:
            pids, rss, backend = bench.sample_tree(os.getpid())
            assert backend == "proc"
            assert os.getpid() in pids
            assert rss.get(os.getpid(), 0) > 0

    def test_ps_backend_matches_proc_tree(self, monkeypatch):
        """macOS 分支在 Linux 上实测：GNU procps 与 BSD ps 的
        `ps -axo pid=,ppid=,rss=` 输出规格一致（无表头、三列 KB）。"""
        if os.name == "nt":
            pytest.skip("POSIX ps backend unavailable on Windows")
        monkeypatch.setenv("T6_SAMPLER", "ps")
        pids, rss, backend = bench.sample_tree(os.getpid())
        assert backend == "ps"
        assert os.getpid() in pids
        assert rss.get(os.getpid(), 0) > 0

    def test_none_backend_degrades_to_single_pid(self, monkeypatch):
        monkeypatch.setenv("T6_SAMPLER", "none")
        pids, rss, backend = bench.sample_tree(os.getpid())
        assert (backend, pids, rss) == ("none", [os.getpid()], {})

    def test_backend_override_wins_over_probe(self, monkeypatch):
        """环境变量覆盖优先于探测（排障口径）。CH-469：原断言写死「无 /proc 时
        默认 proc」，在 macOS 上探测本就返回 ps，等于断言了宿主的文件系统布局。
        改为伪造探测值为 proc，再验覆盖仍是唯一权威——契约与宿主无关。
        """
        monkeypatch.setattr(bench, "_probe_backend", lambda: "proc")
        monkeypatch.setenv("T6_SAMPLER", "ps")
        assert bench._tree_backend() == "ps"
        monkeypatch.delenv("T6_SAMPLER")
        assert bench._tree_backend() == "proc"

    def test_queue_counts_missing_db(self):
        s = bench.Sampler(os.getpid(), Path("/nonexistent/state.sqlite"))
        assert s._queue_counts() is None

    def test_queue_counts_reads_db(self, tmp_path):
        db = tmp_path / "state.sqlite"
        con = sqlite3.connect(db)
        con.execute("CREATE TABLE tasks (job_id TEXT PRIMARY KEY, status TEXT)")
        con.executemany("INSERT INTO tasks VALUES (?, ?)",
                        [("j1", "done"), ("j2", "done"), ("j3", "queued")])
        con.commit()
        con.close()
        s = bench.Sampler(os.getpid(), db)
        assert s._queue_counts() == {"done": 2, "queued": 1}


class TestRunDrainEndToEnd:
    """真 CLI 子进程：doc2md work --enqueue DIR --once --workers N。
    小语料（2 PDF + 1 DOCX + 3 坏文件），只验证仪器接线与状态契约。"""

    def test_tiny_corpus_drain(self, tmp_path):
        corpus = tmp_path / "corpus"
        bench.generate_corpus(corpus, n_pdf=2, n_docx=1, n_bad=3, seed=20261014)
        r = bench.run_drain(corpus, tmp_path / "work", 1)
        # 好文件 3 件全 done；截断 PDF 消费失败（E4151）；notpdf/empty 入队即拒
        assert r["exit"] == 5                      # EXIT_SYS：rep.failed=1
        assert r["final"] == {"done": 3, "failed": 1}
        assert r["enqueue_rejected"] == {"E4151": 2}
        assert r["rejected_codes"] == {"E4151": 1}
        # 峰值积压的上下界（不钉精确值：6 件语料整个 drain 约 0.3s，
        # 0.1s 采样只抓到 1-3 帧，「恰好拍到 4 件同时在场」是时序运气）。
        # 上界 = 真正入队的任务数（3 好件 + 1 截断；notpdf/empty 入队即拒）；
        # 下界取 0（CH-472）：本机 12 轮实测常常只落在**一帧**，而这一帧多半是
        # 终态（queued/running 已清零，backlog_peak=0 占 7/12、=2 占 5/12）。
        # 「采样器接线正常」不由这里证明——它由 test_queue_counts_reads_db 的
        # 确定性断言 + 下面的 queue_series 非空证明；把宿主时序写进下界就是伪覆盖。
        assert 0 <= r["backlog_peak"] <= 4
        # Windows 未装 psutil 时驻留列按设计为 n/a；吞吐/失败账目不受影响。
        if os.name == "nt" and r["sampler"] == "none":
            assert r["peak_tree_mb"] is None and r["rss_available"] is False
        else:
            assert r["peak_tree_mb"] > 0
            # CH-469：sampler 后端断言平台无关（写死 proc 在 macOS 必红）
            assert r["sampler"] in ("proc", "ps", "psutil") and r["rss_available"] is True
        assert r["queue_series"]
        if r["rss_available"]:
            assert r["rss_series"]
        # CH-464：小语料 drain 全程 <0.5s，0.1s 采样可能没有落在半排空区间——
        # half_drain_s 允许 None（排队曲线 queue_series 才是硬口径）
        assert r["half_drain_s"] is None or r["half_drain_s"] >= 0
