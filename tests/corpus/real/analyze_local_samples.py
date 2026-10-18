#!/usr/bin/env python3
"""分析本机收集的真实 Office 模板 / Word 导出 PDF。

样本不在仓库内（Office 模板许可 + 隐私），默认位于 work/real_samples：
  word/*.docx  由 Office 2052 真实模板转换（仅改 content-type）
  pdf/*.pdf    Word 导出（同一模板）
"""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from doc2md.errors import Doc2mdError
from doc2md.pipeline.convert import convert
from doc2md.pipeline.s1_parse import read_pdf

ROOT = Path(__file__).resolve().parents[3] / "work" / "real_samples"
NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CJK = re.compile(r"[\u4e00-\u9fff]")


def _heading_texts(ir) -> list[str]:
    return [b.text.strip() for b in ir.blocks if b.type == "heading" and b.text.strip()]


def _docx_features(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
        names = set(z.namelist())
    feat = {
        "textboxes": xml.count("txbxContent"),
        "numbering": xml.count("<w:numPr"),
        "tables": xml.count("<w:tbl"),
        "merged_cells": xml.count("<w:gridSpan") + xml.count("<w:vMerge"),
        "sdt": xml.count("<w:sdt"),
        "fields": xml.count("<w:fldSimple"),
        "smart_tag": xml.count("<w:smartTag"),
        "hyperlinks": xml.count("<w:hyperlink"),
        "images": xml.count("<w:drawing"),
        "header_parts": len([n for n in names if n.startswith("word/header")]),
        "footer_parts": len([n for n in names if n.startswith("word/footer")]),
    }
    from docx import Document
    d = Document(path)
    feat["paragraphs"] = len(d.paragraphs)
    feat["tables_docx"] = len(d.tables)
    feat["headings_docx"] = sum(
        1 for p in d.paragraphs if p.style and p.style.name.lower().startswith("heading"))
    return feat


def _pdf_metrics(path: Path) -> dict:
    pages = read_pdf(path)
    sizes = Counter()
    cjk = 0
    for page in pages:
        for line in page.lines:
            if line.size:
                sizes[round(line.size, 1)] += max(1, len(line.text.strip()))
            cjk += len(CJK.findall(line.text))
    body = sizes.most_common(1)[0][0] if sizes else None
    return {"pages": len(pages), "body_size": body, "cjk_chars": cjk,
            "size_hist": sizes.most_common(5)}


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def main() -> int:
    word_dir = ROOT / "word"
    pdf_dir = ROOT / "pdf"
    out: dict = {"word": [], "pdf": [], "comparisons": []}
    out_dir = ROOT / "convert_out"
    out_dir.mkdir(parents=True, exist_ok=True)
    for path in sorted(word_dir.glob("*.docx")):
        item = {"file": path.name}
        try:
            item["features"] = _docx_features(path)
        except Exception as exc:  # noqa: BLE001 — 单个模板的包关系损坏不阻断整批
            item["features"] = {"error": f"{type(exc).__name__}: {exc}"}
            item["status"] = f"unreadable:{type(exc).__name__}"
            out["word"].append(item)
            continue
        try:
            r = convert(path, out_dir=out_dir / "docx")
            item["status"] = "ok"
            item["blocks"] = len(r.ir.blocks)
            item["headings"] = _heading_texts(r.ir)
            item["tables"] = sum(1 for b in r.ir.blocks if b.type == "table")
        except Doc2mdError as exc:
            item["status"] = f"reject:{exc.code}"
        except Exception as exc:  # noqa: BLE001
            item["status"] = f"error:{type(exc).__name__}"
        out["word"].append(item)
    for path in sorted(pdf_dir.glob("*.pdf")):
        item = {"file": path.name, "metrics": _pdf_metrics(path)}
        try:
            r = convert(path, out_dir=out_dir / "pdf")
            item["status"] = "ok"
            item["blocks"] = len(r.ir.blocks)
            item["headings"] = _heading_texts(r.ir)
            item["tables"] = sum(1 for b in r.ir.blocks if b.type == "table")
            item["numeric_table_blocks"] = sum(
                1 for b in r.ir.blocks
                if b.type == "table" and re.search(r"\d", b.text or ""))
        except Doc2mdError as exc:
            item["status"] = f"reject:{exc.code}"
        except Exception as exc:  # noqa: BLE001
            item["status"] = f"error:{type(exc).__name__}"
        out["pdf"].append(item)
    by_stem = {Path(w["file"]).stem: w for w in out["word"]}
    for p in out["pdf"]:
        w = by_stem.get(Path(p["file"]).stem)
        if not w or p.get("status") != "ok" or w.get("status") != "ok":
            continue
        truth = {_norm(x) for x in w.get("headings", [])}
        got = {_norm(x) for x in p.get("headings", [])}
        hit = truth & got
        out["comparisons"].append({
            "stem": Path(p["file"]).stem,
            "docx_headings": len(truth),
            "pdf_headings": len(got),
            "matched": len(hit),
            "recall": round(len(hit) / len(truth), 3) if truth else None,
            "precision": round(len(hit) / len(got), 3) if got else None,
            "extra_pdf_headings": sorted(got - truth)[:8],
            "missing_pdf_headings": sorted(truth - got)[:8],
        })
    result = ROOT / "analysis.json"
    result.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"word: {len(out['word'])}  pdf: {len(out['pdf'])}  pairs: {len(out['comparisons'])}")
    bad = [w["file"] for w in out["word"] if w["status"] != "ok"]
    bad += [p["file"] for p in out["pdf"] if p["status"] != "ok"]
    print("non-ok:", bad[:20], "..." if len(bad) > 20 else "")
    for c in out["comparisons"]:
        print(f"{c['stem']}: docx={c['docx_headings']} pdf={c['pdf_headings']} "
              f"match={c['matched']} recall={c['recall']} precision={c['precision']}")
    print(f"report: {result}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
