#!/usr/bin/env python3
"""tests/corpus/make_ideal_docx.py — N-3（QA-OPT-2）：DOCX 标题真值标注生成。

从**源文件 XML 样式**（python-docx 直读 `Heading N`，不经管线任何环节）
程序化生成 `ideal/<样本名>.ideal.json`——与 `expected/` 基线（管线产物）
彻底独立，满足 ideal/README.md 的双口径纪律。provenance 字段如实记录
生成方式：程序化生成、人工可复核，非手写。

用法：`.venv/bin/python tests/corpus/make_ideal_docx.py [--check]`
- 默认：对 sources/docx_*.docx 全量（重）生成；
- `--check`：不写盘，校验现有标注与源样式一致（防标注漂移，供 CI/gate）。

标题真值口径：
- 只取 body 顶层段落（python-docx `Document.paragraphs` 不含表格内段落）；
- 样式名 `Heading N` / `标题 N` → level=N；`Title` → level=1；
- 空文本段落跳过；文本按原样记录（匹配侧自行归一化空白）。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
IDEAL = HERE / "ideal"
STYLE_LEVELS = {"title": 1}


def heading_truth(path: Path) -> list[dict]:
    """源 XML 样式 → 标题真值列表（text/level/seq）。"""
    import docx
    d = docx.Document(str(path))
    out = []
    seq = 0
    for para in d.paragraphs:
        style = (para.style.name or "").strip() if para.style is not None else ""
        level = None
        low = style.lower()
        if low in STYLE_LEVELS:
            level = STYLE_LEVELS[low]
        elif low.startswith("heading ") or low.startswith("标题 "):
            try:
                level = int(style.rsplit(" ", 1)[1])
            except (ValueError, IndexError):
                level = None
        if level is None:
            continue
        text = para.text.strip()
        if not text:
            continue
        seq += 1
        out.append({"seq": seq, "text": text, "level": level})
    return out


def build_annotation(path: Path) -> dict:
    return {
        "source": path.name,
        "kind": "headings",
        "provenance": "程序化生成自源文件 XML 样式（python-docx 直读 Heading N / "
                      "Title，不经管线）；生成脚本 tests/corpus/make_ideal_docx.py，"
                      "人工可复核。与 expected/ 基线（管线产物）无关。",
        "headings": heading_truth(path),
    }


def main() -> int:
    check = "--check" in sys.argv
    sources = sorted(HERE.glob("sources/docx_*.docx"))
    if not sources:
        print("no docx sources", file=sys.stderr)
        return 1
    if not check:
        IDEAL.mkdir(exist_ok=True, parents=True)
    bad = []
    for p in sources:
        ann = build_annotation(p)
        target = IDEAL / f"{p.stem}.ideal.json"
        payload = json.dumps(ann, ensure_ascii=False, indent=2) + "\n"
        if check:
            if not target.is_file():
                bad.append(f"{target.name}: 缺失")
                continue
            if target.read_text(encoding="utf-8") != payload:
                bad.append(f"{target.name}: 与源样式不一致")
        else:
            target.write_text(payload, encoding="utf-8")
            print(f"IDEAL {target.name}: {len(ann['headings'])} headings")
    if check and bad:
        print("\n".join(bad), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
