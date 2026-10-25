#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gen_cases.py —— 测试文档与测试用例的**自动生成器**。

解决的具体问题：现有 `tests/corpus/` 是**合成语料**（reportlab / python-docx
按生成器意图画出来的），它证明了管线「按约定输入」时对，但证明不了
「按真实日常文档」时对。参考资料把 PDF 转换难度明确分档：

    有线框表 ★★ ｜ 跨页表 ★★★★ ｜ 合并单元格 ★★★★ ｜ **无线框表 ★★★★★**
    ｜ 字体编码（CID 无 ToUnicode / Type3 转曲）★★★★★

对照现有 22 份语料，缺口是：**无线框表、跨页表、断词连字符、点线引导符、
数值形态（千分位 / 括号负数）、加密 PDF、页眉逐页变体、目录域**。这些恰恰是
日常文档里最常出事的几类。

本脚本做三件事：
  1. `gen_sources()`  —— 生成上述缺口难度的样本（PDF / DOCX）到 `tests/corpus/hard/`；
  2. `gen_casebook()` —— 按「场景 × 难度」矩阵**自动生成测试用例文档**（Markdown）；
  3. `gen_pytest()`   —— 自动生成 pytest 骨架（断言留 TODO，由人填入口径）。

生成的样本自带 `expectations.json`——**目标输出的结构化声明**（不经过管线），
与 `srcstruct` 的源端直抽互为校验：一个说「我按这个意图画的」，一个说
「我从文件里读出来是这个」。两者不一致即为生成器缺陷，必须先修生成器。

用法：
    .venv/bin/python tests/qa/gen_cases.py --all
    .venv/bin/python tests/qa/gen_cases.py --sources
    .venv/bin/python tests/qa/gen_cases.py --casebook --out docs/qa/casebook.md
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
CORPUS = HERE.parent / "corpus"
HARD = CORPUS / "hard"
REPO = HERE.parents[1]
sys.path.insert(0, str(REPO))

from tests.qa import prioritize  # noqa: E402  # 只读它的 LINK_B / RULE_STAGE 口径

# ============================================================================
# 1. 难度清单（参考资料难度分档 × 本项目非目标边界）
# ============================================================================

#: 每项：难度星级 / 现有覆盖 / 日常出现频率 / 生成函数
DIFFICULTIES: list[dict] = [
    {"id": "D01", "name": "无线框表格", "star": 5, "covered": False, "freq": "高",
     "why": "PDF 不存『第3行第2列』，只存线条与散落文字；无线框只能靠间隙聚类猜，"
            "是最高难度档。现有语料 pdf_03 是有线框，未触达真实难点。",
     "gen": "hard_noframe_table"},
    {"id": "D02", "name": "跨页表格", "star": 4, "covered": False, "freq": "中",
     "why": "表头在上页、数据在下页；跨页续表若不强判会裂成两张表或丢表头。",
     "gen": "hard_cross_page_table"},
    {"id": "D03", "name": "合并单元格（DOCX gridSpan/vMerge）", "star": 4,
     "covered": "partial", "freq": "高",
     "why": "GFM 管道表无 colspan 语义；§1.2 非目标 6 明确不还原，但降级必须诚实留痕。",
     "gen": "hard_merged_cells_docx"},
    {"id": "D04", "name": "断词连字符（de-hyphenation）", "star": 2, "covered": False,
     "freq": "中", "why": "英文行尾 comput-\\ner 需并回 computer，否则正文被切碎。",
     "gen": "hard_hyphenation"},
    {"id": "D05", "name": "点线引导符", "star": 3, "covered": False, "freq": "中",
     "why": "目录/明细表的 ...... 会被当成真实字符，污染文本与表格单元格。",
     "gen": "hard_leader_dots"},
    {"id": "D06", "name": "数值形态（千分位 / 括号负数）", "star": 3, "covered": False,
     "freq": "高", "why": "『(1,234)』被读成正数、千分位逗号丢失——不改位置只改数值，"
                          "比串位更难发现，是数字类文档的头号静默错误。",
     "gen": "hard_numeric_forms"},
    {"id": "D07", "name": "加密 PDF", "star": 1, "covered": False, "freq": "低",
     "why": "E4151 拒收路径；口令加密与结构损坏必须可区分（detail.password_required）。",
     "gen": "hard_encrypted_pdf"},
    {"id": "D08", "name": "页眉逐页变体", "star": 2, "covered": "partial", "freq": "高",
     "why": "O-09 只归一了『第 N 页第 M 栏』；奇偶页眉、章节名页眉仍未覆盖。",
     "gen": "hard_page_header_variants"},
    {"id": "D09", "name": "目录（TOC 域）", "star": 2, "covered": "partial", "freq": "中",
     "why": "TOC 域会把目录文本注入正文；docx_05 有域但未造真 TOC。",
     "gen": "hard_toc_field_docx"},
    {"id": "D10", "name": "CID 字体无 ToUnicode / Type3 转曲", "star": 5,
     "covered": False, "freq": "低",
     "why": "复制出来是 (cid:142) 或乱码。**本项目不作 OCR**，正确行为是门禁拒收或"
            "留 W1046 告警，而非输出乱码——合成成本高，登记为真机样本待办。"
            "本机以**代理腿**覆盖检出路径（`proxy: true`），真实字体样本仍登记待办（B-47）。",
     "gen": "hard_cid_mojibake"},
    {"id": "D11", "name": "双栏阅读顺序", "star": 3, "covered": True, "freq": "高",
     "why": "pdf_10 已覆盖；朴素提取会得到『左一句右一句』乒乓错乱。",
     "gen": None},
    {"id": "D12", "name": "文本框 / SmartArt（DOCX）", "star": 3, "covered": True,
     "freq": "中", "why": "docx_06 已覆盖；文本框内容极易丢或顺序错乱。",
     "gen": None},
]


# ============================================================================
# 2. 缺口难度样本生成器
# ============================================================================

def _new_pdf(name: str):
    """创建 PDF 画布（reportlab，与 make_corpus.py 同口径）。"""
    import reportlab.rl_config
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    reportlab.rl_config.invariant = 1
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    HARD.mkdir(parents=True, exist_ok=True)
    path = HARD / name
    if path.exists():
        path.unlink()
    return path, canvas.Canvas(str(path), pagesize=A4), A4, cm


def hard_noframe_table() -> dict:
    """D01 无线框表格：只靠列间距对齐，无任何线条。"""
    path, c, (W, H), cm = _new_pdf("hard_01_noframe_table.pdf")
    c.setFont("STSong-Light", 16)
    c.drawString(2 * cm, H - 2.5 * cm, "一、无线框表格测试")
    c.setFont("STSong-Light", 10.5)
    rows = [["项目", "本期金额", "上期金额"],
            ["营业收入", "1,234,567", "987,654"],
            ["营业成本", "765,432", "654,321"],
            ["净利润", "469,135", "333,333"]]
    y = H - 4.5 * cm
    cols_x = [2 * cm, 8 * cm, 13 * cm]      # 只靠 x 坐标对齐，无 BOX/GRID
    for r in rows:
        for x, cell in zip(cols_x, r):
            c.drawString(x, y, cell)
        y -= 0.9 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D01", "star": 5,
            "expect": {"tables": [{"nrow": 4, "ncol": 3,
                                   "header": rows[0], "body": rows[1:]}],
                       "headings": [{"text": "一、无线框表格测试", "level": 1}]},
            "judge": "列数与单元格文本须与 expect 一致；无线框是本项唯一难点，"
                     "拿不到 3 列即为缺陷（不是『降级』，是漏检）。"}


def hard_cross_page_table() -> dict:
    """D02 跨页表格：表头在第 1 页，数据续到第 2 页。"""
    path, c, (W, H), cm = _new_pdf("hard_02_cross_page_table.pdf")
    c.setFont("STSong-Light", 16)
    c.drawString(2 * cm, H - 2.5 * cm, "二、跨页表格测试")
    c.setFont("STSong-Light", 10.5)
    header = ["序号", "科目", "金额"]
    cols_x = [2 * cm, 5 * cm, 12 * cm]
    y = H - 4.5 * cm
    for x, cell in zip(cols_x, header):
        c.drawString(x, y, cell)
    y -= 0.8 * cm
    rows = []
    for i in range(1, 41):
        row = [str(i), f"科目{i:02d}", str(1000 + i)]
        rows.append(row)
        for x, cell in zip(cols_x, row):
            c.drawString(x, y, cell)
        y -= 0.62 * cm
        if y < 2.5 * cm and i < 40:
            c.showPage()
            c.setFont("STSong-Light", 10.5)
            y = H - 2.5 * cm
            for x, cell in zip(cols_x, header):   # 续页重画表头（真实文档常这样）
                c.drawString(x, y, cell)
            y -= 0.8 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D02", "star": 4,
            "expect": {"tables": [{"nrow_min": 41, "ncol": 3, "header": header}],
                       "headings": [{"text": "二、跨页表格测试", "level": 1}]},
            "judge": "跨页续表若裂成两张表（n_pred=2）即为缺陷；表头重复出现是否"
                     "合并由人工裁决，但**行数不得丢**。"}


def hard_hyphenation() -> dict:
    """D04 断词连字符：行尾 hyphen 需并回。"""
    path, c, (W, H), cm = _new_pdf("hard_04_hyphenation.pdf")
    c.setFont("STSong-Light", 16)
    c.drawString(2 * cm, H - 2.5 * cm, "四、断词连字符测试")
    c.setFont("Helvetica", 10.5)
    y = H - 4.5 * cm
    # 每行以 hyphen 结尾，下一行接续（模拟真实排版断词）
    for frag in ["The configu-", "ration para-", "meter must be preser-", "ved intact."]:
        c.drawString(2 * cm, y, frag)
        y -= 0.7 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D04", "star": 2,
            "expect": {"contains": ["configuration", "parameter", "preserved"],
                       "headings": [{"text": "四、断词连字符测试", "level": 1}]},
            "judge": "产物须含 configuration / parameter / preserved 三个完整词；"
                     "若出现 'configu- ration' 即为 de-hyphenation 缺失。"}


def hard_leader_dots() -> dict:
    """D05 点线引导符：目录式 ...... 不得污染正文与单元格。"""
    path, c, (W, H), cm = _new_pdf("hard_05_leader_dots.pdf")
    c.setFont("STSong-Light", 16)
    c.drawString(2 * cm, H - 2.5 * cm, "五、点线引导符测试")
    c.setFont("STSong-Light", 10.5)
    y = H - 4.5 * cm
    for item, page in [("第一章 总则", "1"), ("第二章 职责", "12"), ("第三章 附则", "30")]:
        c.drawString(2 * cm, y, item)
        c.drawString(9 * cm, y, "." * 20)
        c.drawString(15.5 * cm, y, page)
        y -= 0.85 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D05", "star": 3,
            "expect": {"headings": [{"text": "五、点线引导符测试", "level": 1}],
                       "contains": ["第一章 总则", "第二章 职责", "第三章 附则"]},
            "judge": "引导符不得成为单元格内容；页码须与条目同行关联。"
                     "若 TOC 形态被误判为表格，须显式降级告警。"}


def hard_numeric_forms() -> dict:
    """D06 数值形态：千分位 / 括号负数 / 百分号 / 单位后缀。"""
    path, c, (W, H), cm = _new_pdf("hard_06_numeric_forms.pdf")
    c.setFont("STSong-Light", 16)
    c.drawString(2 * cm, H - 2.5 * cm, "六、数值形态测试")
    c.setFont("STSong-Light", 10.5)
    y = H - 4.5 * cm
    rows = [["科目", "本期", "上期"],
            ["营业收入", "1,234,567", "987,654"],
            ["营业外支出", "(1,234)", "(5,678)"],
            ["毛利率", "36.5%", "33.8%"],
            ["资产总额", "￥8,765,432", "￥7,654,321"]]
    cols_x = [2 * cm, 8 * cm, 13 * cm]
    for r in rows:
        for x, cell in zip(cols_x, r):
            c.drawString(x, y, cell)
        y -= 0.85 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D06", "star": 3,
            "expect": {"tables": [{"nrow": 5, "ncol": 3,
                                   "header": rows[0], "body": rows[1:]}],
                       "headings": [{"text": "六、数值形态测试", "level": 1}]},
            "judge": "『(1,234)』**必须**保持括号形态——把它读成 1234 或 -1234 都改变了"
                     "数值语义，是数字类文档的头号静默错误。千分位逗号与货币符号须保留。"}


def hard_encrypted_pdf() -> dict:
    """D07 加密 PDF：期望 E4151 拒收，且可区分『口令』与『结构损坏』。"""
    import pypdf
    # 暂存名必须与终名不同：两者同名时 pypdf 的写出目标就是 base 自身，
    # 随后 base.unlink() 会把刚生成的加密样本删掉（D07 资产长期缺失的真因）。
    base, c, (W, H), cm = _new_pdf("hard_07_encrypted.raw.pdf")
    c.setFont("STSong-Light", 12)
    c.drawString(2 * cm, H - 3 * cm, "这是一份口令加密的测试文档")
    c.showPage()
    c.save()
    reader = pypdf.PdfReader(str(base))
    writer = pypdf.PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("doc2md-test-pwd")
    path = HARD / "hard_07_encrypted.pdf"
    with open(path, "wb") as fh:
        writer.write(fh)
    base.unlink(missing_ok=True)
    return {"source": path.name, "difficulty": "D07", "star": 1,
            "expect": {"behavior": "reject:E4151"},
            "judge": "必须拒收且 `detail.password_required` 为真；"
                     "不得与『结构损坏』混淆，也不得静默产出残缺结果（§1.2 非目标 1）。"}


def hard_page_header_variants() -> dict:
    """D08 页眉逐页变体：奇偶页 + 章节名页眉（O-09 只覆盖了『第 N 页第 M 栏』）。"""
    path, c, (W, H), cm = _new_pdf("hard_08_header_variants.pdf")
    chapters = ["总则", "职责分工", "流程规范", "附则"]
    for i in range(8):
        c.setFont("STSong-Light", 9)
        # 逐页变化的页眉：章节名 + 奇偶页差异
        c.drawString(2 * cm, H - 1.5 * cm,
                     f"第 {i // 2 + 1} 章 {chapters[(i // 2) % 4]}　{'奇数页' if i % 2 == 0 else '偶数页'}")
        c.drawString(16 * cm, H - 1.5 * cm, f"- {i + 1} -")
        c.setFont("STSong-Light", 16)
        c.drawString(2 * cm, H - 3.2 * cm, f"第 {i + 1} 节 正文章节标题")
        c.setFont("STSong-Light", 10.5)
        c.drawString(2 * cm, H - 4.5 * cm, "本页正文内容，用于验证页眉剔除不误伤正文。")
        c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D08", "star": 2,
            "expect": {"headings": [{"text": f"第 {i + 1} 节 正文章节标题", "level": 1}
                                    for i in range(8)],
                       "not_contains": ["奇数页", "偶数页"]},
            "judge": "8 个章节标题须全部保留；页眉里的『奇数页/偶数页』与孤立页码"
                     "『- N -』须全部剔除。**误剔除正文是比漏剔除更严重的缺陷**。"}


def _new_docx(name: str):
    """创建 DOCX 画布（与 make_corpus._docx 同口径；末尾必须 _save_docx 冻结）。"""
    from docx import Document

    HARD.mkdir(parents=True, exist_ok=True)
    path = HARD / name
    if path.exists():
        path.unlink()
    return path, Document()


def _save_docx(path: Path, doc) -> None:
    """保存并冻结：python-docx 每次 save 都写当前时刻（zip mtime + core.xml 日期），
    不冻结就没有可复现的样本哈希（CH-442 的同一形态）。"""
    sys.path.insert(0, str(REPO / "tests"))
    from make_golden import freeze_docx

    doc.save(path)
    freeze_docx(path)


def _run(*els):
    from docx.oxml import OxmlElement

    r = OxmlElement("w:r")
    for el in els:
        r.append(el)
    return r


def _fld_char(kind: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    return el


def _instr(text: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = text
    return el


def _text(s: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:t")
    el.set(qn("xml:space"), "preserve")
    el.text = s
    return el


def hard_merged_cells_docx() -> dict:
    """D03 合并单元格：横向 gridSpan + 纵向 vMerge 各一处。"""
    path, d = _new_docx("hard_03_merged_cells.docx")
    d.add_heading("三、合并单元格测试", level=1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    hspan = t.cell(1, 1).merge(t.cell(1, 2))     # gridSpan=2
    vspan = t.cell(2, 0).merge(t.cell(3, 0))     # vMerge
    for (r, c), val in {
        (0, 0): "季度", (0, 1): "收入", (0, 2): "成本",
        (1, 0): "第一季度",
        (2, 1): "2,000,000", (2, 2): "1,500,000",
        (3, 1): "3,000,000", (3, 2): "2,400,000",
    }.items():
        t.cell(r, c).text = val
    hspan.text = "1,234,567 / 987,654"
    vspan.text = "上半年"
    d.add_paragraph("合并单元格的降级形态必须诚实：要么保持 4×3 网格且内容不丢，"
                    "要么降级为 HTML 表格并带 W1045 告警。")
    _save_docx(path, d)
    return {"source": path.name, "difficulty": "D03", "star": 4,
            "expect": {"headings": [{"text": "三、合并单元格测试", "level": 1}],
                       "tables": [{"nrow": 4, "ncol": 3,
                                   "header": ["季度", "收入", "成本"]}],
                       "contains": ["1,234,567 / 987,654", "上半年",
                                    "2,000,000", "3,000,000"]},
            "judge": "§1.2 非目标 6：不还原 colspan/rowspan。产物必须**要么**保持 "
                     "4×3 管道表且 4 个 contains 全中，**要么**降级为 HTML 表格并带 "
                     "W1045；两者都不满足（内容丢失且无告警）即静默降级缺陷。"}


def hard_toc_field_docx() -> dict:
    """D09 目录域：域指令 + **Word 写回的缓存结果**都在文档里。

    真实 .docx 的 TOC 不是「一个占位符」，而是 begin/separate/**缓存目录文本**/end
    四段；缓存文本长得就像正文（含制表位与页码），朴素抽取会把它当正文标题输出，
    于是每个章节标题在产物里出现两次、锚点撞车。

    每个域构件必须**各自成 run**：`w:instrText` 只允许出现在 `w:r` 内（`CT_P` 只接受
    `w:pPr` / `w:hyperlink` / `w:r`），直接挂 `w:p` 是 schema 非法——Word 会修复或
    忽略该指令，按 run 遍历文档的读取器则完全看不见它，样本就不再是「真 TOC 域」。
    缓存条目仍是独立段落（Word 的实际写法），域结束符收在最后一个缓存条目段末。
    """
    path, d = _new_docx("hard_09_toc_field.docx")
    d.add_heading("九、目录域测试", level=1)

    def _para(*els) -> None:
        p = d.add_paragraph()
        for el in els:
            p._p.append(el)

    _para(_run(_fld_char("begin")), _run(_instr(' TOC \\o "1-3" \\h \\z \\u ')),
          _run(_fld_char("separate")))
    entries = [("第一章 总则", "1"), ("第二章 职责分工", "12"), ("第三章 附则", "30")]
    for i, (entry, page) in enumerate(entries):
        els = [_run(_text(f"{entry}\t{page}"))]
        if i == len(entries) - 1:      # Word 把域结束符收在最后一个缓存条目段末
            els.append(_run(_fld_char("end")))
        _para(*els)
    for title, body in [("第一章 总则", "本章规定适用范围与术语定义。"),
                        ("第二章 职责分工", "本章规定各岗位职责与交接流程。"),
                        ("第三章 附则", "本章规定生效日期与解释权归属。")]:
        d.add_heading(title, level=1)
        d.add_paragraph(body)
    _save_docx(path, d)
    return {"source": path.name, "difficulty": "D09", "star": 2,
            "expect": {"headings": [{"text": "九、目录域测试", "level": 1},
                                    {"text": "第一章 总则", "level": 1},
                                    {"text": "第二章 职责分工", "level": 1},
                                    {"text": "第三章 附则", "level": 1}],
                       "contains": ["本章规定适用范围与术语定义。"],
                       "no_duplicate_headings": True},
            "judge": "目录缓存结果不得被当成正文标题重复输出：三个章节标题在产物里"
                     "**各出现一次**（`no_duplicate_headings`）。若产品选择保留目录"
                     "文本，必须与真实章节标题可区分（不产生重复锚点）；引导用的"
                     "制表位与页码不得成为独立段落噪声。"}


def hard_cid_mojibake() -> dict:
    """D10 **代理腿**：文本层就是乱码，验证「检出 / 拒收」路径。

    真实 CID 无 ToUnicode / Type3 转曲需要字体子集化工具，本机不可合成——真机样本
    仍登记为待办（见 §12.7 B-47）。本样本用 Helvetica 把 `(cid:142)` 直接画进文本
    层，等价于「抽取结果就是乱码」这一**下游可观测形态**：它证明的是 S1.5 体检门禁
    能否拦住乱码文本层，**不证明字体解析能力**，故 spec 显式标 `proxy: true`。
    """
    path, c, (W, H), cm = _new_pdf("hard_10_cid_proxy.pdf")
    c.setFont("Helvetica", 12)
    y = H - 3 * cm
    for line in ["(cid:142)(cid:87)(cid:33)(cid:120) (cid:7)",
                 "(cid:12)(cid:99)(cid:7)(cid:200)(cid:31)",
                 "(cid:55) report (cid:66)(cid:77) 2026"]:
        c.drawString(2 * cm, y, line)
        y -= 0.8 * cm
    c.showPage()
    c.save()
    return {"source": path.name, "difficulty": "D10", "star": 5, "proxy": True,
            "expect": {"behavior": "reject:E4221"},
            "judge": "本项目不作 OCR（§1.2 非目标 1），乱码文本层的正确行为是**门禁"
                     "拒收**或**显式告警**，而不是静默产出乱码。契约本身待裁定"
                     "（拒收 E4221 还是新增 W 码，B-47 需人拍板）；裁定前本项以 "
                     "xfail(strict=True) 钉住，不得当作已通过。"}


_GENERATORS = {
    "hard_noframe_table": hard_noframe_table,
    "hard_cross_page_table": hard_cross_page_table,
    "hard_hyphenation": hard_hyphenation,
    "hard_leader_dots": hard_leader_dots,
    "hard_numeric_forms": hard_numeric_forms,
    "hard_encrypted_pdf": hard_encrypted_pdf,
    "hard_page_header_variants": hard_page_header_variants,
    "hard_merged_cells_docx": hard_merged_cells_docx,
    "hard_toc_field_docx": hard_toc_field_docx,
    "hard_cid_mojibake": hard_cid_mojibake,
}


def gen_sources() -> list[dict]:
    """生成全部缺口难度样本，返回 expectations 列表。"""
    out = []
    for d in DIFFICULTIES:
        fn = d.get("gen")
        if not fn or d.get("covered") is True:
            continue
        try:
            spec = _GENERATORS[fn]()
        except Exception as exc:  # noqa: BLE001
            print(f"  SKIP {d['id']} {d['name']}: {type(exc).__name__}: {exc}")
            continue
        spec["name"] = d["name"]
        spec["why"] = d["why"]
        out.append(spec)
        print(f"  GEN  {spec['source']:34s} {d['star']}★ {d['name']}")
    (HARD / "expectations.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


# ============================================================================
# 3. 测试用例文档自动生成（场景 × 难度矩阵）
# ============================================================================

#: 日常使用场景（用户视角）→ 关注指标
SCENARIOS: list[dict] = [
    {"id": "S1", "name": "合同 / 协议归档", "profile": "financial",
     "focus": "条款编号层级、表格金额、跨页续表",
     "key_metrics": ["heading_f1", "table_teds", "order_nid"]},
    {"id": "S2", "name": "财报 / 报表入库", "profile": "financial",
     "focus": "数值形态、合并单元格、表头重复",
     "key_metrics": ["table_teds", "text", "lint_contract"]},
    {"id": "S3", "name": "技术手册 / 文档站", "profile": "manual",
     "focus": "标题大纲、代码块语言、内部链接",
     "key_metrics": ["heading_f1", "structure_bsf", "lint_contract", "asset"]},
    {"id": "S4", "name": "论文 / 研究报告", "profile": "academic",
     "focus": "双栏阅读顺序、脚注尾注、跨页表",
     "key_metrics": ["order_nid", "table_teds", "structure_bsf"]},
    {"id": "S5", "name": "批量历史档案转换", "profile": "batch",
     "focus": "确定性、规范合规、拒收可判定",
     "key_metrics": ["idempotent", "lint_contract", "render", "asset"]},
    {"id": "S6", "name": "扫描件 / 图片型 PDF", "profile": "batch",
     "focus": "门禁拒收而非残缺输出",
     "key_metrics": ["render", "honest"]},
]


# ============================================================================
# 2b. 用例矩阵与卡片的口径表（CH-501 Task 8）
# ============================================================================

#: 场景 → 相关难度。**单一口径**：矩阵表与用例卡都从 `tc_matrix()` 取编号，
#: 加一条难度只改这里；改这里会让 TC 编号整体位移，所以 `--fill-results`
#: 回填时按 (场景, 难度) 对齐、不按编号对齐。
SCENARIO_DIFFICULTY: dict[str, list[str]] = {
    "S1": ["D03", "D08", "D09"],
    "S2": ["D01", "D02", "D03", "D06"],
    "S3": ["D05", "D08", "D09"],
    "S4": ["D02", "D11"],
    "S5": ["D07", "D10"],
    "S6": ["D10"],
}

#: 难度号 → 执行载体（仓库相对路径）。"" = 本机无载体，只能记「未执行」。
CARRIER: dict[str, str] = {
    "D01": "tests/corpus/hard/hard_01_noframe_table.pdf",
    "D02": "tests/corpus/hard/hard_02_cross_page_table.pdf",
    "D03": "tests/corpus/hard/hard_03_merged_cells.docx",
    "D04": "tests/corpus/hard/hard_04_hyphenation.pdf",
    "D05": "tests/corpus/hard/hard_05_leader_dots.pdf",
    "D06": "tests/corpus/hard/hard_06_numeric_forms.pdf",
    "D07": "tests/corpus/hard/hard_07_encrypted.pdf",
    "D08": "tests/corpus/hard/hard_08_header_variants.pdf",
    "D09": "tests/corpus/hard/hard_09_toc_field.docx",
    "D10": "tests/corpus/hard/hard_10_cid_proxy.pdf",   # 代理样本，契约见 B-47
    "D11": "tests/corpus/sources/pdf_10_two_column_6p.pdf",
    "D12": "tests/corpus/sources/docx_06_textbox_image.docx",
}

#: 难度号 → 所属模块（管线阶段，口径同 prioritize.RULE_STAGE，不另写一套）
DIFFICULTY_MODULE: dict[str, str] = {
    "D01": "S1 解析（无线框几何聚类）→ S3 结构",
    "D02": "S1 解析（行装配按 y 聚类）→ S2 分片缝合",
    "D03": "S1 解析（gridSpan/vMerge）→ S4 渲染（GFM 无 colspan）",
    "D04": "S1 解析（run 拼接）→ S3 结构（断词并回）",
    "D05": "S1 解析（run 间空格）→ S3 结构（目录形态识别）",
    "D06": "S1 解析（文本序）→ S4 渲染（数值形态保真）",
    "D07": "S1.5 门禁（加密判定 E4151）",
    "D08": "S1 解析（页眉页脚）→ S3 结构（误升防护）",
    "D09": "S1 解析（TOC 域）→ S3 结构（锚点去重）",
    "D10": "S1.5 门禁（CID / 无 ToUnicode 判定）",
    "D11": "S2 分片缝合（阅读顺序）",
    "D12": "S3 结构（文本框 / 图片资产）",
}

#: 难度号 → 覆盖指标（指标名取自 mdscore / defect_scan.RULES / W 码，不造新名）
DIFFICULTY_METRICS: dict[str, list[str]] = {
    "D01": ["table_teds", "table_shape_ok", "col_count_drift"],
    "D02": ["table_teds", "table_shape_ok", "split_table", "W1045"],
    "D03": ["table_shape_ok", "silent_degrade", "honest"],
    "D04": ["text_ned", "contains"],
    "D05": ["lint_contract", "leader_dots_in_cell", "cjk_space_collapse", "heading_f1"],
    "D06": ["contains", "text_ned"],
    "D07": ["reject:E4151", "password_required"],
    "D08": ["furniture_residue", "heading_f1", "lint_dialect"],
    "D09": ["heading_f1", "duplicate_anchor", "structure_bsf"],
    "D10": ["mojibake_residue", "reject:E4221", "W1046"],
    "D11": ["order_nid"],
    "D12": ["structure_bsf", "asset"],
}

#: 十字段的**唯一**权威顺序（渲染与测试各持一份字面量，顺序与缺失都要被两边看见）
CARD_FIELDS = ("编号", "所属模块", "测试目标", "输入与前置条件", "执行步骤",
               "期望结果", "实际结果", "通过状态", "优先级", "覆盖指标")


def tc_matrix() -> list[dict]:
    """场景 × 难度 → 用例编号（行优先，与矩阵表同一份）。"""
    out, n = [], 0
    for s in SCENARIOS:
        for d in DIFFICULTIES:
            if d["id"] in SCENARIO_DIFFICULTY.get(s["id"], []):
                n += 1
                out.append({"tc": f"TC-{n:03d}", "scenario": s["id"],
                            "scenario_name": s["name"], "difficulty": d["id"],
                            "difficulty_name": d["name"], "star": d["star"],
                            "profile": s["profile"]})
    return out


def _cell(v: object) -> str:
    """Markdown 表格单元格：转义竖线、折叠换行（judge 文本里带 `|` 会把表撞散）。"""
    return str(v).replace("|", "\\|").replace("\n", " ")


def _known_gaps() -> dict[str, str]:
    """难度号 → 登记 B 号。

    口径**只取自 `tests/qa/test_hard_gate.KNOWN_GAPS`**（缺口 ↔ B 号的唯一切换点）：
    在卡片里再抄一份，就会出现「xfail 行已删、卡片仍写 B-43」的两处口径分家。
    按路径加载而非 import：那是 pytest 测试模块，不该被生成器长期持有。
    """
    import importlib.util

    path = HERE / "test_hard_gate.py"
    spec = importlib.util.spec_from_file_location("d2m_hard_gate_cards", path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"无法加载 {path}（缺口↔B 号口径缺失，卡片不能凭空判 B 号）")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return dict(mod.KNOWN_GAPS)


def fill_results(results_dir: Path) -> dict:
    """读 work/qa/*.json。**缺文件就记 missing，不猜、不沿用旧值。**"""
    ev: dict = {"dir": str(results_dir), "missing": [], "generated_at": {}, "data": {},
                "known_gaps": {}}
    for name in ("realtest", "attribution", "defects", "priorities", "eval", "hard"):
        p = results_dir / f"{name}.json"
        if not p.is_file():
            ev["missing"].append(name)
            continue
        blob = json.loads(p.read_text(encoding="utf-8"))
        ev["data"][name] = blob
        ev["generated_at"][name] = blob.get("generated_at", "?")
    if "hard" in ev["data"]:
        ev["known_gaps"] = _known_gaps()
    return ev


def _source_record(ev: dict, layer: str, source: str) -> dict | None:
    for r in (ev["data"].get(layer, {}) or {}).get("records", []):
        if r.get("source") == source:
            return r
    return None


def _hard_record(ev: dict, difficulty: str) -> dict | None:
    for r in (ev["data"].get("hard", {}) or {}).get("records", []):
        if r.get("difficulty") == difficulty:
            return r
    return None


def card_status(difficulty: str, ev: dict) -> tuple[str, str]:
    """(通过状态, 实际结果)。四值封闭：通过 / 失败 / 缺口·B-xx / 未执行。"""
    carrier = CARRIER.get(difficulty, "")
    if not carrier:
        return "未执行", "本机无载体样本（需真机文档，见 §12.7 登记）"
    if not ev["data"]:
        return "未执行", f"本轮未提供证据（缺 {', '.join(ev['missing']) or '全部'}）"
    src = Path(carrier).name
    # 难度样本（D01–D10）不在 attribution / defects 的 22 篇语料清单里，它们的
    # 本轮实测值只能来自难度腿自己的 JSON（`run_hard.py --json`）。
    hard = _hard_record(ev, difficulty)
    if hard is not None:
        stamp = ev["generated_at"].get("hard", "?")
        bad = hard.get("failed_rows") or []
        detail = (f"结构面 {'全绿' if hard.get('ok') else f'{len(bad)} 项未达'}；"
                  f"未达项 {_cell(json.dumps(bad, ensure_ascii=False))[:400]}；"
                  f"证据轮次 generated_at={stamp}")
        if hard.get("ok"):
            return "通过", detail
        b = ev["known_gaps"].get(difficulty, "")
        return (f"缺口 · {b}" if b else "失败"), detail
    att = _source_record(ev, "attribution", src)
    dsc = _source_record(ev, "defects", src)
    rt = ev["data"].get("realtest", {})
    ts = {c["name"]: c for ly in rt.get("layers", []) for c in ly.get("checks", [])}
    stamp = ev["generated_at"].get("attribution", "?")
    if att is None and dsc is None:
        why = f"证据里没有 {src} 的记录"
        if difficulty in ("D01", "D02", "D03", "D04", "D05", "D06", "D07", "D08",
                          "D09", "D10"):
            why += "（难度腿不在 attribution 的 22 篇语料清单里，需 run_hard 输出）"
        return "未执行", f"{why}；run_hard 状态见 realtest.json 的 L5.hard_leg"
    blocking = [d for d in (att or {}).get("deviations", [])
                if d.get("severity") == "blocking"]
    hits = [f for f in (dsc or {}).get("findings", [])
            if f.get("severity") in ("blocking", "major")]
    b = next((prioritize.LINK_B.get(f["rule"], "") for f in hits
              if prioritize.LINK_B.get(f["rule"])), "")
    detail = (f"total_gap={(att or {}).get('total_gap', '?')}；"
              f"blocking 偏差 {len(blocking)} 条；规则命中 "
              f"{_cell(json.dumps((dsc or {}).get('by_rule', {}), ensure_ascii=False))}；"
              f"证据轮次 generated_at={stamp}")
    if blocking or hits:
        if b:
            return f"缺口 · {b}", detail
        return "失败", detail
    hard_leg = ts.get("hard_leg")
    if hard_leg and hard_leg["status"] != "pass":
        return "未执行", f"难度腿未跑通：{hard_leg['detail']}"
    return "通过", detail


def _release_blocking() -> set[str]:
    """§12.7 标了「发布阻断」的 B 号集合（派生自文档，不手抄名单）。"""
    dp = REPO / "docs" / "dev-plan.md"
    if not dp.is_file():
        return set()
    return prioritize.release_blocking(dp.read_text(encoding="utf-8"))


def card_priority(status: str, ev: dict) -> str:
    """优先级：**由分档派生**，不手填。缺口/失败 → 查 priorities.json 里同题项的档；
    查不到 → P1（有实测偏差但没进优先级表，本身就是要看的异常）。通过 → P2（守成）。

    但 §12.7 标了「发布阻断」的 B 号**无条件 P0**，优先于 priorities.json 的查找：
    难度腿的缺口不在 22 篇语料的评分输入里，只走下面的查表会把 B-44 记成 P1，
    而口径「阻断即 P0、不看修复成本」正是本轮定的（`prioritize` 同一份文档派生）。
    """
    if status == "通过":
        return "P2（守成：本卡已由自动化覆盖，只需保持绿）"
    b = status.split("·")[-1].strip() if "·" in status else ""
    if b and b in _release_blocking():
        return "P0（§12.7 发布阻断：阻断即 P0，不看修复成本、不等 priorities.json 收编）"
    for it in (ev["data"].get("priorities", {}) or {}).get("items", []):
        if b and it.get("b_linked") == b:
            return f"{it['priority']}（阻断={it['blocking']}，分数={it['score']}）"
    return "P1（有偏差但未进优先级表——先查 priorities.json 是否缺输入层）"


def gen_cards(ev: dict) -> list[str]:
    specs = {}
    exp = HARD / "expectations.json"
    if exp.is_file():
        specs = {s["difficulty"]: s for s in json.loads(exp.read_text(encoding="utf-8"))}
    man = {}
    mp = CORPUS / "manifest.json"
    if mp.is_file():
        man = {Path(e["source"]).name: e
               for e in json.loads(mp.read_text(encoding="utf-8"))}
    L: list[str] = ["## 五、用例卡（十字段）", "",
                    "> 十字段口径见 `gen_cases.CARD_FIELDS`。**期望结果只来自声明侧**"
                    "（`expectations.json` / `manifest.json`），实际结果只来自本轮 "
                    "`work/qa/*.json` 回填并带 `generated_at`；没跑的一律写「未执行」+ 原因。",
                    ""]
    if ev.get("missing"):
        L += [f"> ⚠️ 本轮缺证据层：{', '.join(ev['missing'])}——相关卡片的「实际结果」"
              "只能记「未执行」，不许沿用旧值。", ""]
    for e in tc_matrix():
        d = e["difficulty"]
        status, actual = card_status(d, ev)
        prio = card_priority(status, ev)
        carrier = CARRIER.get(d, "")
        spec = specs.get(d)
        if spec:
            want = (f"{_cell(spec['judge'])}｜声明：`{_cell(json.dumps(spec['expect'], ensure_ascii=False)[:220])}`"
                    f"（出处 `tests/corpus/hard/expectations.json`，声明侧不经管线）")
        else:
            m = man.get(Path(carrier).name, {})
            want = (f"`expected_behavior={m.get('expected_behavior', '?')}`，"
                    f"逐字节等于 `{m.get('expected_markdown', '?')}`"
                    f"（出处 `tests/corpus/manifest.json`）")
        steps = ("1) `.venv/bin/python -m tests.qa.realtest --only L5`"
                 "  2) `.venv/bin/python tests/corpus/run_hard.py "
                 f"{d}`（或 `run_corpus.py --only {Path(carrier).stem}`）"
                 "  3) 读 `work/qa/{attribution,defects,priorities}.json` 中该源记录")
        L += [f"### {e['tc']} ｜ {e['scenario']} {e['scenario_name']} × "
              f"{d} {e['difficulty_name']}（{'★' * e['star']}，画像 `{e['profile']}`）",
              "", "| 字段 | 内容 |", "|---|---|",
              f"| 编号 | {e['tc']} |",
              f"| 所属模块 | {DIFFICULTY_MODULE[d]} |",
              f"| 测试目标 | {_cell(spec['judge'][:60]) if spec else '与已入库基线逐字节一致'} |",
              f"| 输入与前置条件 | `{carrier or '（本机无载体）'}`；离线、"
              "`llm.mode=offline`、默认配置、`--force` |",
              f"| 执行步骤 | {steps} |",
              f"| 期望结果 | {want} |",
              f"| 实际结果 | {_cell(actual)} |",
              f"| 通过状态 | {status} |",
              f"| 优先级 | {_cell(prio)} |",
              f"| 覆盖指标 | {', '.join(DIFFICULTY_METRICS[d])} |", ""]
    return L


def gen_casebook(*, results_dir: Path | None = None) -> str:
    """生成测试用例文档 Markdown。"""
    L = ["# Doc2md 测试用例集（自动生成）", "",
         "> 由 `tests/qa/gen_cases.py` 生成。**难度清单**来自 PDF 转换的公开难度分档，"
         "**场景清单**来自日常实际使用；两者交叉处即为必须覆盖的用例。",
         "> 指标口径见 `tests/qa/mdscore.py`，真值口径见 `tests/qa/srcstruct.py`。", "",
         "## 一、难度清单（D-xx）", "",
         "| ID | 难度项 | 星级 | 日常频率 | 现有覆盖 | 为什么难 / 边界 |",
         "|---|---|---|---|---|---|"]
    for d in DIFFICULTIES:
        cov = {True: "✅ 已覆盖", False: "❌ **缺口**"}.get(d["covered"], "🟡 部分")
        L.append(f"| {d['id']} | {d['name']} | {'★' * d['star']} | {d['freq']} | "
                 f"{cov} | {d['why']} |")

    L += ["", "## 二、使用场景（S-xx）", "",
          "| ID | 场景 | 画像 | 关注点 | 关键指标 |", "|---|---|---|---|---|"]
    for s in SCENARIOS:
        L.append(f"| {s['id']} | {s['name']} | `{s['profile']}` | {s['focus']} | "
                 f"{', '.join(s['key_metrics'])} |")

    L += ["", "## 三、用例矩阵（场景 × 难度 → 用例编号 TC-xx）", "",
          "> 生成规则：场景关注点与难度项相关即产出一条用例；"
          "「—」表示该组合不构成有效用例（如扫描件场景不测双栏）。", "",
          "| 场景 \\ 难度 | " + " | ".join(d["id"] for d in DIFFICULTIES) + " |",
          "|" + "---|" * (len(DIFFICULTIES) + 1)]

    cells: dict[tuple[str, str], str] = {}
    for e in tc_matrix():
        cells[(e["scenario"], e["difficulty"])] = e["tc"]
    for s in SCENARIOS:
        row = [f"**{s['id']} {s['name']}**"]
        for d in DIFFICULTIES:
            row.append(cells.get((s["id"], d["id"]), "—"))
        L.append("| " + " | ".join(row) + " |")

    L += ["", f"共生成 **{len(cells)}** 条交叉用例。", "",
          "## 四、缺口难度的目标声明（expectations.json）", "",
          "> 由 `gen_sources()` 随样本一同生成，**不经过管线**，是「我按这个意图画的」；"
          "与 `srcstruct` 的「我从文件里读出来是这个」互为校验。", ""]
    exp_path = HARD / "expectations.json"
    if exp_path.is_file():
        for e in json.loads(exp_path.read_text(encoding="utf-8")):
            L.append(f"### {e['difficulty']} {e['name']}（{'★' * e['star']}）—— `{e['source']}`")
            L.append("")
            L.append(f"- 判定口径：{e['judge']}")
            L.append(f"- 目标声明：`{json.dumps(e['expect'], ensure_ascii=False)}`")
            L.append("")
    else:
        L.append("_（尚未生成样本，先跑 `gen_cases.py --sources`）_")

    ev = (fill_results(results_dir) if results_dir
          else {"missing": [], "data": {}, "generated_at": {}, "known_gaps": {}})
    L += gen_cards(ev)
    return "\n".join(L) + "\n"


# ============================================================================
# 4. pytest 骨架自动生成
# ============================================================================

def gen_pytest(out: Path) -> None:
    """生成 pytest 骨架：断言留 TODO，由人按判定口径填值。"""
    L = ['"""难度语料回归（自动生成骨架）。', '',
         '由 `tests/qa/gen_cases.py --pytest` 生成。断言值需人工按 expectations.json',
         '的 `judge` 口径确认后填入——**不得**直接把当前产物当期望值写死',
         '（那是 C-31 的自啮合陷阱：只能证明稳定，不能证明对）。', '"""',
         'from __future__ import annotations', '', 'import json', 'import sys',
         'from pathlib import Path', '', 'import pytest', '',
         'HERE = Path(__file__).resolve().parent',
         'CORPUS = HERE.parent / "corpus"',
         'HARD = CORPUS / "hard"',
         'sys.path.insert(0, str(HERE.parents[1]))', '',
         'SPECS = json.loads((HARD / "expectations.json").read_text(encoding="utf-8")) '
         'if (HARD / "expectations.json").is_file() else []', '',
         '@pytest.mark.corpus',
         '@pytest.mark.parametrize("spec", SPECS, ids=lambda s: s["source"])',
         'def test_hard_corpus(spec, tmp_path):',
         '    """逐难度样本：跑转换并对账目标声明。"""',
         '    from doc2md.errors import Doc2mdError',
         '    from doc2md.pipeline.convert import convert',
         '    src = HARD / spec["source"]',
         '    if not src.is_file():',
         '        pytest.skip(f"{spec[\'source\']} 未生成")',
         '    try:',
         '        res = convert(src, out_dir=tmp_path, force=True)',
         '    except Doc2mdError as exc:',
         '        want = spec["expect"].get("behavior")',
         '        assert want == f"reject:{exc.code}", \\',
         '            f"{spec[\'source\']}: 期望 {want}，实际 {exc.code}"',
         '        return',
         '    md = res.markdown',
         '    exp = spec["expect"]',
         '    gaps: list[str] = []',
         '    for token in exp.get("contains", []):',
         '        if token not in md:',
         '            gaps.append(f"缺关键内容 {token!r}")',
         '    for token in exp.get("not_contains", []):',
         '        if token in md:',
         '            gaps.append(f"残留噪声 {token!r}")',
         '    if gaps:',
         '        # 已知缺陷走 xfail（不转红、但 CI 输出可见），口径见',
         '        # docs/qa/eval-report-v2.md；缺陷修复后应立即转为红断言。',
         '        pytest.xfail(f"{spec[\'difficulty\']} 已知缺陷: " + "; ".join(gaps))',
         '    pytest.skip(f"TODO 未填断言: {spec[\'difficulty\']} —— {spec[\'judge\']}")',
         ]
    out.write_text("\n".join(L) + "\n", encoding="utf-8")


def _raw_text(path: Path) -> str:
    """源端直读（**不 import doc2md**）：自检必须独立于被测管线，否则等于自证。"""
    if path.suffix.lower() == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    from docx import Document

    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _toc_field_gaps(path: Path) -> list[str]:
    """D09 源端结构核对：域标记 begin/separate/end 齐全 + `w:instrText` 确是 TOC 指令。

    直接读 `word/document.xml`（zip 直读，不经 python-docx 段落 API）：域标记本身
    不产生段落文本，`_raw_text` 看不见它们——TOC 域被删掉时内容层面完全隐形，这
    正是本检查存在的理由。`w:instrText` 还必须挂在 `w:r` 下（schema 要求）：挂在
    `w:p` 下时 Word 会修复或忽略该指令，按 run 遍历文档的读取器则完全看不见它。
    """
    import re as _re
    import zipfile
    from xml.etree import ElementTree as ET

    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    gaps = [f'缺 w:fldChar w:fldCharType="{kind}"'
            for kind in ("begin", "separate", "end")
            if not _re.search(rf'<w:fldChar[^>]*w:fldCharType="{kind}"', xml)]
    instr = _re.findall(r"<w:instrText[^>]*>(.*?)</w:instrText>", xml, _re.S)
    if not any("TOC" in t for t in instr):
        gaps.append('缺含 "TOC" 的 w:instrText')
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = ET.fromstring(xml)
    parents = {child: parent for parent in root.iter() for child in parent}
    if any(el.tag == f"{w}instrText" and parents[el].tag != f"{w}r"
           for el in root.iter()):
        gaps.append("w:instrText 未挂在 w:r 内（schema 非法，run 级读取器看不见）")
    return gaps


def _load_expectations() -> list[dict]:
    """读 expectations.json；缺失 / 为空 / 不可解析一律返回空表。

    由 `--selfcheck` 的 `if not rows` 守卫转成非零退出——原实现缺省读文件后打印
    `SELFCHK 0/0` 并退出 0，等于「没有样本」被当成「全部通过」。
    """
    path = HARD / "expectations.json"
    if not path.is_file():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []
    return data if isinstance(data, list) else []


def selfcheck_sources() -> list[tuple[str, bool, str]]:
    """生成器自检：expectations 的构造意图 ↔ 源端直读。

    两者不一致 = **生成器画错了**，必须先修生成器再谈产品缺陷；否则「产品没通过」
    可能只是「样本没画对」（第十七轮 O-g / G9 的同一形态：真值取自产物或与产物
    同源，就看不见损失）。自检只核生成器能控制的东西：
      1. `contains` 的每个 token 确实在源文件里（制表位按空白归一后比对；
         行尾断词 token 按「连字符 + 换行」并回后比对，D04 即此形态）；
      2. `tables[i]` 的 ncol / nrow 与源端直抽一致（DOCX 侧）；PDF 无线框表
         没有线框可抽（srcstruct 能力边界），退到「声明的单元格内容 token
         都在源文本里」，几何不可判定这一事实写进检查说明；
      3. `behavior: reject:*` 的样本不核内容（加密 / 乱码样本本就不该有可用文本），
         但标了 `proxy` 的代理腿（D10）要核「文本层确实是 `(cid:` 乱码形态」——
         换成普通 ASCII 时，「门禁拒收乱码」这个样本前提就没了；
      4. `headings` 的每条文本确实在源端原文里（D08 只有 headings / not_contains，
         旧口径从不读它，却报「构造意图与源端一致」＝空转）；
      5. 目录域样本（`no_duplicate_headings`）的 begin/separate/end 与 TOC 指令
         确实在 `word/document.xml` 里（域标记不产生段落文本，3 核不到）。
    """
    import re as _re

    from tests.qa import srcstruct

    specs = _load_expectations()
    rows: list[tuple[str, bool, str]] = []
    for s in specs:
        did, src = s["difficulty"], HARD / s["source"]
        if not src.is_file():
            rows.append((did, False, f"{s['source']} 不存在"))
            continue
        exp = s.get("expect", {})
        bad: list[str] = []
        notes: list[str] = []
        if exp.get("behavior"):
            if s.get("proxy"):
                try:
                    has_cid = "(cid:" in _raw_text(src)
                except Exception as exc:  # noqa: BLE001
                    rows.append((did, False,
                                 f"源端直读失败 {type(exc).__name__}: {exc}"))
                    continue
                if not has_cid:
                    rows.append((did, False,
                                 "代理腿文本层无 (cid: token（乱码形态没画进去）"))
                    continue
                notes.append("代理腿乱码文本层含 (cid: token")
            rows.append((did, True, "; ".join(notes) or "拒收样本，跳过内容自检"))
            continue
        norm = lambda t: _re.sub(r"\s+", "", t or "")   # noqa: E731
        try:
            text = _raw_text(src)
        except Exception as exc:  # noqa: BLE001
            rows.append((did, False, f"源端直读失败 {type(exc).__name__}: {exc}"))
            continue
        raw = norm(text)
        if exp.get("no_duplicate_headings") and src.suffix.lower() == ".docx":
            field_gaps = _toc_field_gaps(src)
            if field_gaps:
                bad.extend(field_gaps)
            else:
                notes.append("TOC 域 begin/separate/end + instrText 源端核对")
        dehyph = norm(_re.sub(r"-\s*\n\s*", "", text))
        gaps, used_dehyph = [], False
        for tok in exp.get("contains", []):
            if norm(tok) in raw:
                continue
            if norm(tok) in dehyph:
                used_dehyph = True
                continue
            gaps.append(tok)
        if gaps:
            rows.append((did, False, f"生成器没画进去：{gaps}"))
            continue
        if used_dehyph:
            notes.append("断词 token 按行尾连字符并回核对")
        heads = exp.get("headings", [])
        if heads:
            missing_heads = [h.get("text") for h in heads
                             if norm(h.get("text")) not in raw]
            if missing_heads:
                bad.append(f"headings 源端缺 {missing_heads}")
            else:
                notes.append(f"headings 源端核对（{len(heads)} 条）")
        tables = exp.get("tables", [])
        if tables and src.suffix.lower() == ".docx":
            gt = srcstruct.extract(src)
            table_bad = 0
            for i, want in enumerate(tables):
                if i >= len(gt.get("tables", [])):
                    bad.append(f"table[{i}] 源端直抽只有 {len(gt.get('tables', []))} 张")
                    table_bad += 1
                    continue
                got = gt["tables"][i]
                if "ncol" in want and got.get("ncol") != want["ncol"]:
                    bad.append(f"table[{i}] ncol 源端 {got.get('ncol')} != 声明 {want['ncol']}")
                    table_bad += 1
                if "nrow" in want and got.get("nrow") != want["nrow"]:
                    bad.append(f"table[{i}] nrow 源端 {got.get('nrow')} != 声明 {want['nrow']}")
                    table_bad += 1
            if not table_bad:
                notes.append(f"table 几何源端核对（{len(tables)} 张）")
        elif tables:
            # srcstruct 的 PDF 表格清单只在有线框时可信；无线框表改核内容 token，
            # 否则 D01 / D02 / D06 会被源端读取器的能力边界误报成生成器缺陷。
            for i, want in enumerate(tables):
                cells = list(want.get("header", [])) + [c for row in want.get("body", [])
                                                        for c in row]
                missing = [c for c in cells if c and norm(c) not in raw]
                if missing:
                    bad.append(f"table[{i}] 源端缺内容 {missing}")
            notes.append("PDF 无线框表按内容 token 核对（几何不可判定）")
        rows.append((did, not bad,
                     "; ".join(bad) or "; ".join(notes) or "构造意图与源端一致"))
    return rows


# ============================================================================

def main() -> int:
    ap = argparse.ArgumentParser(description="Doc2md 测试文档与用例自动生成器")
    ap.add_argument("--all", action="store_true", help="样本 + 用例文档 + pytest 骨架")
    ap.add_argument("--sources", action="store_true", help="生成缺口难度样本")
    ap.add_argument("--casebook", action="store_true", help="生成测试用例文档")
    ap.add_argument("--selfcheck", action="store_true",
                    help="生成器自检：expectations 的构造意图 ↔ 源端直读（不经管线）")
    ap.add_argument("--out", default=None, help="用例文档输出路径")
    ap.add_argument("--fill-results", default="", metavar="DIR",
                    help="从该目录的 work/qa/*.json 回填十字段卡片的实际结果/通过状态/优先级")
    ap.add_argument("--pytest", default=None, help="pytest 骨架输出路径")
    args = ap.parse_args()
    if args.fill_results:
        args.casebook = True
    if not (args.all or args.sources or args.casebook or args.pytest or args.selfcheck):
        args.all = True

    if args.all or args.sources:
        print("生成缺口难度样本 ->", HARD)
        gen_sources()

    if args.all or args.casebook:
        out = Path(args.out) if args.out else (HERE.parents[1] / "docs" / "qa" / "casebook.md")
        rd = Path(args.fill_results) if args.fill_results else None
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(gen_casebook(results_dir=rd), encoding="utf-8")
        print(f"CASEBOOK {out} cards={len(tc_matrix())} "
              f"fill={'yes:' + str(rd) if rd else 'no'}")

    if args.all or args.pytest:
        out = Path(args.pytest) if args.pytest else (HERE / "test_hard_corpus.py")
        gen_pytest(out)
        print("pytest 骨架 ->", out)

    if args.selfcheck:
        rows = selfcheck_sources()
        for did, ok, msg in rows:
            print(f"  {'OK  ' if ok else 'FAIL'} {did}: {msg}")
        if not rows:
            # 没有样本 ≠ 全部通过：缺失 / 空 / 不可解析的 expectations.json 必须
            # 显式失败，否则 --selfcheck 会以 `SELFCHK 0/0` + 退出 0 假绿。
            print("SELFCHK 0/0 —— expectations.json 缺失、为空或无法解析；"
                  "先跑 .venv/bin/python tests/qa/gen_cases.py --sources",
                  file=sys.stderr)
            return 1
        bad = [r for r in rows if not r[1]]
        print(f"SELFCHK {len(rows) - len(bad)}/{len(rows)}")
        return 1 if bad else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
