"""tests/test_qa_o05_cell_br.py — O-05 单元格 <br> + 降级诚实标记（QA-D2M-2026-10-21 / C-06/07/29）。

缺陷：GFM 表达力天花板的三处降维对用户不可见——
① 单元格内多段落以单空格相连（段落边界不可辨，审阅面看成一句）；
② 合并单元格（gridSpan / vMerge）拍平为网格空位，无任何信号；
③ 跨片表格不合并等降级只在 meta.warnings 里沉底，报告与 API 不显式暴露。

修复（F-11「诚实暴露」路线，不做自动合并）：
- S1 `_docx_cell_text` 段落边界改 `<br>`（GFM 单元格内合法换行）；
- S3 `_gfm_table` 统一收口单元格规范（`|` 转义、换行统一 `<br>`）；
- S1 `Line.row_flat` → S3 表级 `flags.flattened_cells` → convert 留
  W1045 告警；`render_report` 顶部出提示行（iff warnings 非空）；
- `GET /docs/<id>` 响应透出 `warnings` 计数。
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from doc2md.errors import WARN_TABLE_FLATTENED
from doc2md.pipeline.s3_structure import _gfm_table

pytestmark = [pytest.mark.milestone("DoD")]
CORPUS = Path(__file__).parent / "corpus"


# ---------------------------------------------------------------- 单元：_gfm_table
def test_gfm_table_escapes_pipe_and_maps_breaks():
    """单元格规范统一收口：`|` 转义（否则提前结束单元格）、换行统一 `<br>`
    （S1 段落边界标记与 run 内 w:br 映射的 \\n 同一口径）、首尾空白剥掉。"""
    md = _gfm_table([["列一", "列|二"], ["a\nb", "  padded  "]])
    assert "\\|" in md and md.count("|") > 4
    assert "a<br>b" in md
    assert "padded" in md and "  padded  " not in md


def test_gfm_table_does_not_escape_br_itself():
    """`<br>` 是 GFM 单元格内的合法换行，不得被转义成字面量。"""
    md = _gfm_table([["甲", "乙"], ["一段<br>二段", "丙"]])
    assert "一段<br>二段" in md


# ---------------------------------------------------------------- 单元：拍平计数
def test_merged_cell_docx_gets_flat_count_and_warning(tmp_path):
    """gridSpan / vMerge 拍平 → 表级 flags.flattened_cells 非零 →
    W1045 告警 + 报告提示行（诚实暴露三件套在一条链上）。"""
    import docx

    from doc2md.pipeline.convert import convert

    d = docx.Document()
    d.add_heading("合并单元格文档", level=1)
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "表头"
    t.cell(0, 1).text = "甲"
    t.cell(0, 2).text = "乙"
    t.cell(1, 0).merge(t.cell(1, 2))          # 横向合并：gridSpan=3
    t.cell(1, 0).text = "跨三列的内容"
    p = tmp_path / "merged.docx"
    d.save(str(p))

    r = convert(p, out_dir=tmp_path / "out")
    tables = [b for b in r.ir.blocks if b.type == "table"]
    assert tables, "合并单元格样本应产出表格块"
    flat = sum((b.flags or {}).get("flattened_cells", 0) for b in tables)
    assert flat >= 2, f"gridSpan 展开空位应被计数，实际 {flat}"
    assert WARN_TABLE_FLATTENED in r.meta["warnings"]
    with zipfile.ZipFile(r.zip_path) as z:
        report = z.read(f"{r.doc_id}/report.html").decode("utf-8")
    assert 'class="degraded"' in report, "有告警时报告必须出顶部提示行"
    assert WARN_TABLE_FLATTENED in report


def test_plain_table_has_no_flat_count_no_warning_no_notice(tmp_path):
    """反例（防噪音红线）：无合并的常规表格——零拍平计数、零 W1045、
    报告零提示；单段落单元格产物与改动前逐字节一致（G6）。

    样本带一段长正文：S1.5 门禁对每页字符数有 120 下限（chars<120 即
    E4221 拒收），纯小表格样本过不了门禁——那是门禁语义，不是本测试的面。
    """
    import docx

    from doc2md.pipeline.convert import convert

    d = docx.Document()
    d.add_paragraph("常规表格样本的正文部分，用于越过门禁的每页字符下限，"
                    "内容本身与表格断言无关，长度足够即可顺利进入管线。"
                    "再补一段确保稳妥，避免样本漂移到门禁边缘。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "字段"
    t.cell(0, 1).text = "值"
    t.cell(1, 0).text = "甲"
    t.cell(1, 1).text = "1"
    p = tmp_path / "plain.docx"
    d.save(str(p))

    r = convert(p, out_dir=tmp_path / "out")
    tables = [b for b in r.ir.blocks if b.type == "table"]
    assert tables and all(not (b.flags or {}).get("flattened_cells") for b in tables)
    assert WARN_TABLE_FLATTENED not in r.meta["warnings"]
    with zipfile.ZipFile(r.zip_path) as z:
        report = z.read(f"{r.doc_id}/report.html").decode("utf-8")
    assert 'class="degraded"' not in report, "无告警时报告不得出提示行（防恒真）"
    assert "| 字段 | 值 |" in r.markdown


# ---------------------------------------------------------------- gate：iff 双向断言
def test_gate_report_notice_iff_warnings():
    """O-05 门禁：报告含提示 iff ir.warnings 非空（双向断言，防恒真/防噪音）。

    对全部语料样本实转：有告警的样本报告必含提示行；无告警的样本必不含。
    """
    import json
    import tempfile

    from doc2md.pipeline.convert import convert

    entries = json.loads((CORPUS / "manifest.json").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="o05-gate-") as td:
        out = Path(td) / "out"
        out.mkdir()
        checked = 0
        for e in entries:
            if e["expected_behavior"] != "ok":
                continue
            src = CORPUS / "sources" / e["source"]
            r = convert(src, out_dir=out)
            checked += 1
            warns = list(r.meta.get("warnings") or [])
            with zipfile.ZipFile(r.zip_path) as z:
                report = z.read(f"{r.doc_id}/report.html").decode("utf-8")
            has_notice = 'class="degraded"' in report
            assert has_notice == bool(warns), (
                f"{e['source']}: 提示行与 warnings 不同构 "
                f"(notice={has_notice}, warnings={warns})")
        assert checked >= 20, f"覆盖样本不足：{checked}"


# ---------------------------------------------------------------- API：warnings 透出
def test_docs_response_surfaces_warnings(saved_doc):
    """O-05：GET /docs/<id> 响应含 warnings 列表（调用方不必翻 meta 嵌套）。"""
    client, tokens, doc_id = saved_doc
    r = client.get(f"/docs/{doc_id}", headers={"X-Doc2md-Token": tokens["reader"]})
    assert r.status_code == 200
    body = r.get_json()
    assert isinstance(body.get("warnings"), list)
