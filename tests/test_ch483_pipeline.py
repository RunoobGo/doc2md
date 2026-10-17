"""CH-483 / 管线功能完整性：标题、DOCX 结构、marks、raw 名。"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from doc2md.pipeline.s1_parse import Line, RawPage, read_docx
from doc2md.pipeline.s3_structure import analyze

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = f'xmlns:w="{W}"'
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _lines(path: Path) -> list[Line]:
    return [ln for pg in read_docx(path) for ln in pg.lines]


def _add_raw(parent, xml: str):
    from docx.oxml import parse_xml
    el = parse_xml(xml)
    parent.append(el)
    return el


# --- R8-PIPE-1：中文空格章标题 -------------------------------------------------
@pytest.mark.gate
def test_cjk_spaced_chapter_heading_detected():
    ir = analyze([RawPage(index=0, lines=[
        Line("第 1 章 总则", size=18.0, bold=True),
        Line("正文。", size=10.5),
    ])])
    h = ir.blocks[0]
    assert h.type == "heading" and h.level == 1, h


@pytest.mark.gate
def test_golden_320p_has_all_chapter_headings(large_single):
    heads = [b for b in large_single.ir.blocks if b.type == "heading"]
    assert sum(1 for b in heads if b.level == 1) == 64
    assert len(heads) == 640, len(heads)


def test_cjk_spaced_chapter_long_sentence_is_not_heading():
    ir = analyze([RawPage(index=0, lines=[
        Line("第 3 章所述内容，详见后文。", size=10.5),
    ])])
    assert ir.blocks[0].type == "paragraph"


# --- R8-PIPE-3：auto 噪声判据 ---------------------------------------------------
@pytest.mark.gate
def test_auto_gate_rejects_noise_only_pages():
    from doc2md.errors import E4221_GateReject
    from doc2md.pipeline.s1_5_gate import check
    pages = [RawPage(index=i, lines=[Line("###", size=10.0) for _ in range(10)])
             for i in range(2)]
    with pytest.raises(E4221_GateReject) as ei:
        check(pages, min_chars=10, mode="auto")
    page_reasons = [r for p in ei.value.detail["page_detail"]
                    for r in p["reasons"]]
    assert any("noise" in r for r in page_reasons), ei.value.detail


# --- R8-PIPE-2：内容控件 / customXml -------------------------------------------
@pytest.mark.gate
def test_block_sdt_and_custom_xml_are_extracted(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("普通段落")
    _add_raw(d.element.body, (
        f'<w:sdt {NS}><w:sdtContent>'
        '<w:p><w:r><w:t>内容控件内正文</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'))
    _add_raw(d.element.body, (
        f'<w:customXml {NS}>'
        '<w:p><w:r><w:t>自定义XML内正文</w:t></w:r></w:p>'
        '</w:customXml>'))
    p = tmp_path / "sdt.docx"
    d.save(str(p))
    texts = [ln.text for ln in _lines(p)]
    assert texts == ["普通段落", "内容控件内正文", "自定义XML内正文"], texts


@pytest.mark.gate
def test_inline_sdt_text_is_not_dropped(tmp_path):
    import docx
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("前")
    _add_raw(p._p, (
        f'<w:sdt {NS}><w:sdtContent>'
        '<w:r><w:t>中</w:t></w:r>'
        '</w:sdtContent></w:sdt>'))
    p.add_run("后")
    path = tmp_path / "inline_sdt.docx"
    d.save(str(path))
    assert _lines(path)[0].text == "前中后"


# --- R8-PIPE-4：自动编号 --------------------------------------------------------
_NUMBERING = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering {NS}>
  <w:abstractNum w:abstractNumId="0">
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1."/></w:lvl>
    <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>
      <w:lvlText w:val="%2)"/></w:lvl>
  </w:abstractNum>
  <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
</w:numbering>
"""


def _inject_numbering(path: Path, paragraphs: list[tuple[int, str]]) -> None:
    """给 docx 加 numbering part，并追加 (ilvl, text) 自动编号段落。"""
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    body = "".join(
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{lvl}"/>'
        f'<w:numId w:val="1"/></w:numPr></w:pPr>'
        f'<w:r><w:t>{text}</w:t></w:r></w:p>'
        for lvl, text in paragraphs)
    doc = items["word/document.xml"].decode("utf-8")
    items["word/document.xml"] = doc.replace(
        "</w:body>", body + "</w:body>").encode("utf-8")
    # python-docx 默认模板通常已带 numbering.xml + 关系；只在缺失时补关系，
    # 否则同一 reltype 出现两条会让 part_related_by 抛 ValueError。
    had_numbering = "word/numbering.xml" in items
    items["word/numbering.xml"] = _NUMBERING.encode("utf-8")
    if not had_numbering:
        rels = items["word/_rels/document.xml.rels"].decode("utf-8")
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rIdNum1" Type="{R_NS}/numbering" '
            'Target="numbering.xml"/></Relationships>')
        items["word/_rels/document.xml.rels"] = rels.encode("utf-8")
        ct = items["[Content_Types].xml"].decode("utf-8")
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/word/numbering.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
            "</Types>")
        items["[Content_Types].xml"] = ct.encode("utf-8")
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items():
            z.writestr(name, data)
    tmp.replace(path)


@pytest.mark.gate
def test_docx_auto_numbering_becomes_list_items(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("正文段。")
    path = tmp_path / "numbering.docx"
    d.save(str(path))
    _inject_numbering(path, [(0, "第一项"), (0, "第二项"), (1, "子项")])
    lines = _lines(path)
    numbered = [ln for ln in lines if ln.list_marker]
    assert [(ln.list_level, ln.list_marker) for ln in numbered] == [
        (0, "1."), (0, "2."), (1, "a)")], numbered
    ir = analyze([RawPage(index=0, lines=lines)])
    items = [b.text for b in ir.blocks if b.type == "list_item"]
    assert items[:2] == ["1. 第一项", "2. 第二项"], items
    assert items[2].endswith("子项") and items[2].startswith("  "), items[2]


# --- R8-PIPE-5：合并单元格 ------------------------------------------------------
@pytest.mark.gate
def test_grid_span_expands_to_table_grid(tmp_path):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "跨两列"
    t.cell(1, 0).text = "A2"
    t.cell(1, 1).text = "B2"
    tc = t.rows[0].cells[0]._tc
    # 先移除被合并的第二格，再加 gridSpan：python-docx 的 row.cells 在
    # gridSpan 已存在时会把同一 tc 返回两次，后删会把锚点格一起删掉。
    t.rows[0]._tr.remove(t.rows[0].cells[1]._tc)
    gs = OxmlElement("w:gridSpan")
    gs.set(qn("w:val"), "2")
    tc.get_or_add_tcPr().append(gs)
    path = tmp_path / "grid.docx"
    d.save(str(path))
    rows = [ln.row for ln in _lines(path) if ln.row]
    assert rows == [["跨两列", ""], ["A2", "B2"]], rows


# --- R8-PIPE-6：单元格 br/tab ---------------------------------------------------
def test_cell_keeps_breaks_and_tabs(tmp_path):
    import docx
    from docx.oxml import OxmlElement
    d = docx.Document()
    t = d.add_table(rows=1, cols=1)
    p = t.cell(0, 0).paragraphs[0]
    r = p.add_run("A")
    r._r.append(OxmlElement("w:br"))
    r2 = p.add_run("B")
    r2._r.append(OxmlElement("w:tab"))
    p.add_run("C")
    path = tmp_path / "cell.docx"
    d.save(str(path))
    row = next(ln for ln in _lines(path) if ln.row)
    assert row.row == ["A\nB\tC"], row.row


# --- R8-PIPE-7：marks 生产 ------------------------------------------------------
@pytest.mark.gate
def test_uniform_marks_are_produced_but_mixed_are_not(tmp_path):
    import docx
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("全段粗体").bold = True
    p.add_run("也粗。").bold = True
    p2 = d.add_paragraph()
    p2.add_run("局部粗体").bold = True
    p2.add_run("普通文字。")
    path = tmp_path / "marks.docx"
    d.save(str(path))
    ir = analyze(read_docx(path))
    paras = [b for b in ir.blocks if b.type == "paragraph"]
    assert paras[0].marks.get("bold") is True, paras[0].marks
    assert not paras[1].marks.get("bold"), paras[1].marks


# --- R8-PIPE-8：raw 名不进 zip --------------------------------------------------
@pytest.mark.gate
def test_source_name_raw_not_in_zip_ir_json(golden_dir, cfg, tmp_path):
    from doc2md.pipeline.convert import convert
    res = convert(golden_dir / "docx" / "sample_simple.docx", cfg=cfg,
                  out_dir=tmp_path, raw_name="../../evil<b>.pdf")
    with zipfile.ZipFile(res.zip_path) as z:
        ir_raw = z.read(f"{res.doc_id}/ir.json").decode("utf-8")
        report = z.read(f"{res.doc_id}/report.html").decode("utf-8")
    meta = json.loads(ir_raw)["meta"]
    assert "source_name_raw" not in meta, meta
    assert "../../evil<b>" not in ir_raw
    assert "evil&lt;b&gt;" in report, "report.html 仍应转义展示原始名"
